# ==========================================
# Flask Imports
# ==========================================

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    session,
    flash,
    send_file
)

# ==========================================
# Security
# ==========================================

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

import secrets

# ==========================================
# Database
# ==========================================

import sqlite3

# ==========================================
# Date & Time
# ==========================================

from datetime import datetime, timedelta
import time

# ==========================================
# File Handling
# ==========================================

import os
import io
from io import BytesIO

# ==========================================
# Excel Handling
# ==========================================

import pandas as pd
from openpyxl import load_workbook

# ==========================================
# Random
# ==========================================

import random

# ==========================================
# HTTP Requests
# ==========================================

import requests

# ==========================================
# Environment Variables
# ==========================================

from dotenv import load_dotenv

# ==========================================
# Email
# ==========================================

from flask_mail import Mail, Message

# ==========================================
# QR Code
# ==========================================

import qrcode

# ==========================================
# PDF Generation
# ==========================================

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image
)



app = Flask(__name__)
app.secret_key = "prepnova_super_secret_key"

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# =====================================
# EMAIL CONFIGURATION
# =====================================

app.config["MAIL_SERVER"] = "smtp.gmail.com"
app.config["MAIL_PORT"] = 587
app.config["MAIL_USE_TLS"] = True
app.config["MAIL_USE_SSL"] = False

# Your PrepNova Gmail
app.config["MAIL_USERNAME"] = "prepnovacbt@gmail.com"

# Your Gmail App Password
app.config["MAIL_PASSWORD"] = "qbgh egdx ifuw glhu"

app.config["MAIL_DEFAULT_SENDER"] = (
    "PrepNova CBT",
    "prepnovacbt@gmail.com"
)

mail = Mail(app)

# =====================================
# SEND EMAIL HELPER
# =====================================

def send_email(subject, recipient, body):

    try:

        msg = Message(
            subject=subject,
            recipients=[recipient]
        )

        msg.body = body

        mail.send(msg)

        return True

    except Exception as e:

        print("Email Error:", e)

        return False

app.secret_key = os.getenv("SECRET_KEY")

ADMIN_EMAIL = os.getenv("ADMIN_EMAIL")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD")

PAYSTACK_PUBLIC_KEY = os.getenv("PAYSTACK_PUBLIC_KEY")
PAYSTACK_SECRET_KEY = os.getenv("PAYSTACK_SECRET_KEY")

APP_URL = os.getenv("APP_URL")


# ✅ JAMB course combinations
JAMB_COURSES = {

    "Medicine & Surgery": [
        "Use of English",
        "Biology",
        "Chemistry",
        "Physics"
    ],

    "Nursing": [
        "Use of English",
        "Biology",
        "Chemistry",
        "Physics"
    ],

    "Pharmacy": [
        "Use of English",
        "Biology",
        "Chemistry",
        "Physics"
    ],

    "Computer Science": [
        "Use of English",
        "Mathematics",
        "Physics",
        "Chemistry"
    ],

    "Mechanical Engineering": [
        "Use of English",
        "Mathematics",
        "Physics",
        "Chemistry"
    ],

    "Civil Engineering": [
        "Use of English",
        "Mathematics",
        "Physics",
        "Chemistry"
    ],

    "Law": [
        "Use of English",
        "Literature in English",
        "Government",
        "Christian Religious Studies"
    ],

    "Law (Islamic Studies)": [
        "Use of English",
        "Literature in English",
        "Government",
        "Islamic Religious Studies"
    ],

    "Mass Communication": [
        "Use of English",
        "Literature in English",
        "Government",
        "Christian Religious Studies"
    ],

    "Mass Communication (IRS)": [
        "Use of English",
        "Literature in English",
        "Government",
        "Islamic Religious Studies"
    ],

    "Accounting": [
        "Use of English",
        "Mathematics",
        "Economics",
        "Commerce"
    ],

    "Business Administration": [
        "Use of English",
        "Mathematics",
        "Economics",
        "Commerce"
    ],

    "Economics": [
        "Use of English",
        "Mathematics",
        "Economics",
        "Government"
    ]
}

WAEC_SUBJECTS = [

    "English",
    "Mathematics",
    "Biology",
    "Chemistry",
    "Physics",
    "Economics",
    "Government",
    "Literature",
    "CRS",
    "IRS",
    "Commerce",
    "Accounting",
    "Geography",
    "Civic Education",
    "Agricultural Science",
    "Further Mathematics",
    "Computer Studies"

]

# ✅ Create database and tables
def init_db():
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Users table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        email TEXT,
        password TEXT
    )
    """)

    # Questions table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS questions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        exam_type TEXT,
        subject TEXT,
        question_text TEXT,
        option_a TEXT,
        option_b TEXT,
        option_c TEXT,
        option_d TEXT,
        correct_answer TEXT,
        explanation TEXT
    )
    """)
    
    # ==============================
    # Bookmarked Questions
    # ==============================

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS bookmarks
        (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            username TEXT NOT NULL,

            question_id INTEGER NOT NULL,

            exam_type TEXT NOT NULL,

            subject TEXT NOT NULL,

            bookmarked_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            UNIQUE(username, question_id)
        )
        """
    )
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS result_subjects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        result_id TEXT NOT NULL,
        username TEXT NOT NULL,
        exam_type TEXT NOT NULL,
        exam_name TEXT,
        subject TEXT NOT NULL,
        score INTEGER DEFAULT 0,
        total_questions INTEGER DEFAULT 0,
        percentage REAL DEFAULT 0,
        time_spent INTEGER DEFAULT 0
    )
    """)

    conn.commit()
    conn.close()


# Run once when app starts
init_db()



from datetime import datetime, timedelta

def is_session_valid():

    if "user" not in session:
        return False, "not_logged_in"

    if "session_token" not in session:
        return False, "missing_session"

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            last_activity
        FROM user_sessions
        WHERE
            username=?
            AND session_token=?
            AND is_active=1
        """,
        (
            session["user"],
            session["session_token"]
        )
    )

    row = cursor.fetchone()

    if not row:

        conn.close()

        return False, "another_device"

    session_id = row[0]
    last_activity = row[1]

    if last_activity:

        last_activity = datetime.strptime(
            last_activity,
            "%Y-%m-%d %H:%M:%S"
        )

        if datetime.utcnow() - last_activity > timedelta(minutes=30):

            cursor.execute(
                """
                UPDATE user_sessions
                SET is_active=0
                WHERE id=?
                """,
                (session_id,)
            )

            conn.commit()
            conn.close()

            return False, "timeout"

    conn.close()

    return True, None


@app.before_request
def validate_user_session():

    # Routes that do NOT require login
    allowed_routes = [
        "login",
        "register",
        "admin_login",
        "static",
        "home"
    ]

    if request.endpoint in allowed_routes:
        return

    # Check only student sessions
    if "user" in session:

        valid, reason = is_session_valid()

        if not valid:

            session.clear()

            if reason == "timeout":

                return redirect("/login?session_expired=1")

            elif reason == "another_device":

                return redirect("/login?logged_out=another_device")

            return redirect("/login")

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE user_sessions
            SET last_activity = CURRENT_TIMESTAMP
            WHERE username = ?
            AND session_token = ?
            """,
            (
                session["user"],
                session["session_token"]
            )
        )

        conn.commit()
        conn.close()
        
        
        
def can_take_exam(username):

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            is_active,
            end_date
        FROM subscriptions
        WHERE username=?
        """,
        (username,)
    )

    subscription = cursor.fetchone()

    conn.close()

    if not subscription:
        return False

    is_active = subscription[0]
    end_date = subscription[1]

    if not is_active:
        return False

    from datetime import datetime

    if end_date:

        expiry = datetime.strptime(
            end_date,
            "%Y-%m-%d %H:%M:%S"
        )

        if expiry >= datetime.now():
            return True

    return False



@app.route("/create_login_attempts_table")
def create_login_attempts_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS login_attempts (

        id INTEGER PRIMARY KEY AUTOINCREMENT,

        email TEXT NOT NULL,

        ip_address TEXT NOT NULL,

        attempt_time TEXT DEFAULT CURRENT_TIMESTAMP,

        successful INTEGER DEFAULT 0

    )
    """)

    conn.commit()
    conn.close()

    return "login_attempts table created successfully."


@app.route("/update_subscriptions_table")
def run_update_subscriptions_table():

    update_subscriptions_table()

    return "Subscriptions table updated successfully."


@app.before_request
def check_active_session():

    # Ignore users who are not logged in
    if "user" not in session:
        return

    # Ignore admin for now
    if "admin" in session:
        return

    session_token = session.get("session_token")

    if not session_token:
        return

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT is_active
        FROM user_sessions
        WHERE session_token=?
        """,
        (
            session_token,
        )
    )

    row = cursor.fetchone()

    conn.close()

    # Session no longer active
    if not row or row[0] == 0:

        session.clear()

        return redirect("/login?logged_out=another_device")



# ✅ Home route
@app.route("/")
def home():
    return render_template("index.html")


# ✅ Register route
@app.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        name = request.form["name"].strip()

        email = request.form["email"].strip().lower()

        phone = request.form.get("phone", "").strip()

        password = request.form["password"]

        confirm_password = request.form["confirm_password"]

        if password != confirm_password:
            return "Passwords do not match"

        password = generate_password_hash(password)

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        # Check if email already exists
        cursor.execute(
            """
            SELECT *
            FROM users
            WHERE email=?
            """,
            (email,)
        )

        existing_user = cursor.fetchone()

        if existing_user:
            conn.close()
            return "Email already registered"

        # -----------------------------------
        # Registration Date & Free Trial
        # -----------------------------------

        from datetime import datetime, timedelta

        start_date = datetime.now()

        end_date = start_date + timedelta(days=7)

        date_joined = start_date.strftime("%Y-%m-%d %H:%M:%S")

        # -----------------------------------
        # Create User
        # -----------------------------------

        cursor.execute(
            """
            INSERT INTO users
            (
                name,
                email,
                phone,
                password,
                date_joined
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                name,
                email,
                phone,
                password,
                date_joined
            )
        )

        # -----------------------------------
        # Create 7-Day Free Trial
        # -----------------------------------

        cursor.execute(
            """
            INSERT INTO subscriptions
            (
                username,
                plan_id,
                plan_name,
                start_date,
                end_date,
                payment_reference,
                payment_status,
                is_active,
                amount_paid,
                currency
            )
            VALUES
            (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                email,
                0,
                "Free Trial",
                start_date.strftime("%Y-%m-%d %H:%M:%S"),
                end_date.strftime("%Y-%m-%d %H:%M:%S"),
                None,
                "FREE_TRIAL",
                1,
                0,
                "NGN"
            )
        )

        # =====================================
        # Generate Email Verification Token
        # =====================================

        token = secrets.token_urlsafe(32)

        expires_at = datetime.utcnow() + timedelta(hours=24)

        cursor.execute(
            """
            INSERT INTO email_verification_tokens
            (
                email,
                token,
                expires_at
            )
            VALUES
            (?, ?, ?)
            """,
            (
                email,
                token,
                expires_at
            )
        )

        conn.commit()

        conn.close()

        # =====================================
        # Send Verification Email
        # =====================================

        verification_link = (
            f"http://127.0.0.1:5000/verify_email/{token}"
        )

        email_body = f"""
Hello {name},

Welcome to PrepNova CBT!

Thank you for creating your account.

Please verify your email address by clicking the link below:

{verification_link}

This verification link will expire in 24 hours.

If you did not create this account, please ignore this email.

Regards,

PrepNova CBT Team
"""

        send_email(
            subject="Verify Your PrepNova CBT Email",
            recipient=email,
            body=email_body
        )

        return redirect("/login")

    return render_template("register.html")


@app.route("/create_subscriptions_table")
def create_subscriptions_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS subscriptions (

        id INTEGER PRIMARY KEY AUTOINCREMENT,

        email TEXT NOT NULL,

        plan TEXT NOT NULL,

        start_date TEXT,

        expiry_date TEXT,

        payment_reference TEXT UNIQUE,

        payment_status TEXT DEFAULT 'PENDING',

        is_active INTEGER DEFAULT 0,

        created_at TEXT DEFAULT CURRENT_TIMESTAMP

    )
    """)

    conn.commit()
    conn.close()

    return "subscriptions table created successfully."



@app.route("/update_subscriptions_table")
def update_subscriptions_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    try:
        cursor.execute(
            "ALTER TABLE subscriptions ADD COLUMN payment_provider TEXT"
        )
    except:
        pass

    try:
        cursor.execute(
            "ALTER TABLE subscriptions ADD COLUMN amount_paid REAL DEFAULT 0"
        )
    except:
        pass

    try:
        cursor.execute(
            "ALTER TABLE subscriptions ADD COLUMN currency TEXT DEFAULT 'NGN'"
        )
    except:
        pass

    conn.commit()
    conn.close()

    return "subscriptions table updated successfully."



@app.route("/create_free_trials_table")
def create_free_trials_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS free_trials (

        id INTEGER PRIMARY KEY AUTOINCREMENT,

        email TEXT NOT NULL,

        exam_type TEXT NOT NULL,

        used INTEGER DEFAULT 0,

        used_at TEXT

    )
    """)

    conn.commit()
    conn.close()

    return "free_trials table created successfully."



@app.route("/create_subscription_plans_table")
def create_subscription_plans_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS subscription_plans (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            plan_name TEXT UNIQUE,

            price REAL,

            duration_days INTEGER,

            description TEXT,

            is_active INTEGER DEFAULT 1,

            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP

        )
    """)

    conn.commit()
    conn.close()

    return "subscription_plans table created successfully."



@app.route("/seed_subscription_plans")
def seed_subscription_plans():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    plans = [

        (
            "Monthly",
            5000,
            30,
            "30 Days Unlimited CBT Access"
        ),

        (
            "Quarterly",
            12000,
            90,
            "90 Days Unlimited CBT Access"
        ),

        (
            "Yearly",
            45000,
            365,
            "365 Days Unlimited CBT Access"
        )

    ]

    for plan in plans:

        cursor.execute(
            """
            INSERT OR IGNORE INTO subscription_plans
            (
                plan_name,
                price,
                duration_days,
                description
            )
            VALUES (?, ?, ?, ?)
            """,
            plan
        )

    conn.commit()
    conn.close()

    return "Subscription plans added successfully."



@app.route("/update_payments_table")
def update_payments_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    columns = [

        (
            "paystack_reference",
            "TEXT"
        ),

        (
            "gateway_response",
            "TEXT"
        ),

        (
            "currency",
            "TEXT DEFAULT 'NGN'"
        ),

        (
            "verified_at",
            "TIMESTAMP"
        )

    ]

    for column_name, column_type in columns:

        try:

            cursor.execute(
                f"""
                ALTER TABLE payments
                ADD COLUMN {column_name} {column_type}
                """
            )

        except sqlite3.OperationalError:
            # Column already exists
            pass

    conn.commit()
    conn.close()

    return "Payments table updated successfully."



@app.route("/check_subscription")
def check_subscription():

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute("""
            SELECT *
            FROM subscriptions
            WHERE username=?
        """, (session["user"],))

        data = cursor.fetchone()

        conn.close()

        return str(data)
    

@app.route("/test_email")
def test_email():

    try:

        msg = Message(
            subject="PrepNova CBT Email Test",
            recipients=["aduagba90@gmail.com"]
        )

        msg.body = """
Hello,

Congratulations!

Your PrepNova CBT email system is working successfully.

This email was sent from your Flask application.

Regards,

PrepNova CBT
"""

        mail.send(msg)

        return "✅ Test email sent successfully!"

    except Exception as e:

        return f"❌ {e}"



# ✅ Login route
@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form["email"].strip().lower()
        password = request.form["password"]
        ip_address = request.remote_addr
        user_agent = request.headers.get("User-Agent")

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        # ----------------------------------------------------
        # Check failed login attempts in the last 15 minutes
        # ----------------------------------------------------
        cursor.execute(
            """
            SELECT COUNT(*)
            FROM login_attempts
            WHERE
                (
                    email = ?
                    OR ip_address = ?
                )
            AND successful = 0
            AND attempt_time >= datetime('now', '-15 minutes')
            """,
            (
                email,
                ip_address
            )
        )

        failed_attempts = cursor.fetchone()[0]

        if failed_attempts >= 5:

            conn.close()

            return """
            <h2>Too many failed login attempts.</h2>
            <p>Please wait 15 minutes before trying again.</p>
            """

        # ----------------------------------------------------
        # Check user credentials
        # ----------------------------------------------------
        cursor.execute(
            """
            SELECT *
            FROM users
            WHERE email=?
            """,
            (email,)
        )

        user = cursor.fetchone()

        if user and check_password_hash(user[2], password):
            
            # ------------------------------------
            # Check Email Verification
            # ------------------------------------

            if user[9] == 0:

                conn.close()

                return render_template(
                    "verify_email_required.html"
                )

            # Account suspended
            if user[7] == 0:

                conn.close()

                return render_template(
                    "account_suspended.html"
                )

            # Record successful login
            cursor.execute(
                """
                INSERT INTO login_attempts
                (
                    email,
                    ip_address,
                    successful
                )
                VALUES
                (?, ?, 1)
                """,
                (
                    email,
                    ip_address
                )
            )

            # Clear previous failed attempts
            cursor.execute(
                """
                DELETE FROM login_attempts
                WHERE email=?
                AND successful=0
                """,
                (email,)
            )

            # Generate secure session token
            session_token = secrets.token_hex(32)

            # Deactivate previous sessions
            cursor.execute(
                """
                UPDATE user_sessions
                SET is_active=0
                WHERE username=?
                """,
                (email,)
            )

            # Save new session
            cursor.execute(
                """
                INSERT INTO user_sessions
                (
                    username,
                    session_token,
                    login_time,
                    last_activity,
                    is_active,
                    user_agent,
                    ip_address
                )
                VALUES
                (?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP, 1, ?, ?)
                """,
                (
                    email,
                    session_token,
                    user_agent,
                    ip_address
                )
            )

            conn.commit()
            conn.close()
            
            session.clear()

            session["user"] = email
            session["name"] = user[3]
            session["session_token"] = session_token

            return redirect("/dashboard")

        # ----------------------------------------------------
        # Invalid login
        # ----------------------------------------------------
        cursor.execute(
            """
            INSERT INTO login_attempts
            (
                email,
                ip_address,
                successful
            )
            VALUES
            (?, ?, 0)
            """,
            (
                email,
                ip_address
            )
        )

        conn.commit()
        conn.close()

        return render_template(
            "login.html",
            login_error="Invalid email or password.",
            password_changed=False,
            logged_out=False,
            session_expired=False
        )

    password_changed = (
        request.args.get("password_changed") == "1"
    )
    
    logged_out = (
        request.args.get("logged_out") == "another_device"
    )
    
    session_expired = (
        request.args.get("session_expired") == "1"
    )

    return render_template(
        "login.html",
        password_changed=password_changed,
        logged_out=logged_out,
        session_expired=session_expired,
        login_error=None
    )


# ✅ Active Login Sessions
@app.route("/active_sessions")
def active_sessions():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            session_token,
            login_time,
            last_activity,
            user_agent,
            ip_address
        FROM user_sessions
        WHERE
            username=?
            AND is_active=1
        ORDER BY login_time DESC
        """,
        (
            session["user"],
        )
    )

    sessions = cursor.fetchall()

    conn.close()

    return render_template(
        "active_sessions.html",
        sessions=sessions,
        current_session=session.get("session_token")
    )


# ✅ Logout a specific session
@app.route("/logout_session/<int:session_id>")
def logout_session(session_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get the session the user wants to log out
    cursor.execute(
        """
        SELECT
            session_token,
            username
        FROM user_sessions
        WHERE id=?
        """,
        (
            session_id,
        )
    )

    row = cursor.fetchone()

    # Session doesn't exist
    if not row:

        conn.close()

        return redirect("/active_sessions")

    session_token = row[0]
    username = row[1]

    # Prevent users from logging out someone else's session
    if username != session["user"]:

        conn.close()

        return redirect("/active_sessions")

    # Prevent logging out the current session
    if session_token == session.get("session_token"):

        conn.close()

        return redirect("/active_sessions")

    # Mark the selected session as inactive
    cursor.execute(
        """
        UPDATE user_sessions
        SET is_active=0
        WHERE id=?
        """,
        (
            session_id,
        )
    )

    conn.commit()
    conn.close()

    return redirect("/active_sessions")



# ==========================================
# Forgot Password
# ==========================================
@app.route("/forgot_password", methods=["GET", "POST"])
def forgot_password():

    if request.method == "POST":

        email = request.form["email"].strip().lower()

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        # Check if email exists
        cursor.execute(
            """
            SELECT id
            FROM users
            WHERE email=?
            """,
            (email,)
        )

        user = cursor.fetchone()

        if not user:

            conn.close()

            return "No account found with this email."

        # Disable all previous unused reset links
        cursor.execute(
            """
            UPDATE password_reset_tokens
            SET used=1
            WHERE email=?
            AND used=0
            """,
            (email,)
        )

        # Generate secure token
        token = secrets.token_urlsafe(32)

        # Expire after 30 minutes
        expires_at = datetime.utcnow() + timedelta(minutes=30)

        # Save new token
        cursor.execute(
            """
            INSERT INTO password_reset_tokens
            (
                email,
                token,
                expires_at
            )
            VALUES
            (?, ?, ?)
            """,
            (
                email,
                token,
                expires_at
            )
        )

        conn.commit()

        conn.close()

        # Create reset link
        reset_link = f"http://127.0.0.1:5000/reset_password/{token}"

        msg = Message(
            subject="Reset Your PrepNova CBT Password",
            recipients=[email]
        )

        msg.body = f"""
Hello,

We received a request to reset your PrepNova CBT password.

Click the link below to create a new password:

{reset_link}

This link will expire in 30 minutes.

If you did not request this, simply ignore this email.

Regards,

PrepNova CBT
"""

        mail.send(msg)

        return "Password reset email sent successfully."

    return render_template("forgot_password.html")



# ==========================================
# Reset Password
# ==========================================
@app.route("/reset_password/<token>", methods=["GET", "POST"])
def reset_password(token):

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Check if token exists
    cursor.execute(
        """
        SELECT
            email,
            expires_at,
            used
        FROM password_reset_tokens
        WHERE token=?
        """,
        (token,)
    )

    token_data = cursor.fetchone()

    if not token_data:

        conn.close()

        return render_template(
            "invalid_reset_link.html"
        )

    email = token_data[0]
    expires_at = token_data[1]
    used = token_data[2]

    # Convert database datetime string to Python datetime
    expires_at = datetime.fromisoformat(expires_at)

    # Check if token has expired
    if datetime.utcnow() > expires_at:

        conn.close()

        return render_template(
            "expired_reset_link.html"
        )

    # Check if token has already been used
    if used == 1:

        conn.close()

        return """
        <h2>Reset Link Already Used</h2>
        <p>This password reset link has already been used.</p>
        <a href="/forgot_password">Request another reset link</a>
        """

    # User submitted new password
    if request.method == "POST":

        password = request.form["password"]
        confirm_password = request.form["confirm_password"]

        if password != confirm_password:

            conn.close()

            return """
            <h2>Password Mismatch</h2>
            <p>The passwords you entered do not match.</p>
            <a href="">Go Back</a>
            """
        confirm_password = request.form["confirm_password"]

        if password != confirm_password:

            conn.close()

            return """
            <h2>Password Mismatch</h2>
            <p>The passwords you entered do not match.</p>
            <a href="javascript:history.back()">Go Back</a>
            """

        hashed_password = generate_password_hash(password)

        # Update password
        cursor.execute(
            """
            UPDATE users
            SET password=?
            WHERE email=?
            """,
            (
                hashed_password,
                email
            )
        )

        # Mark token as used
        cursor.execute(
            """
            UPDATE password_reset_tokens
            SET used=1
            WHERE token=?
            """,
            (token,)
        )

        conn.commit()

        conn.close()

        return render_template(
            "password_reset_success.html"
        )

    conn.close()

    return render_template(
        "reset_password.html",
        token=token
    )



@app.route("/verify_email/<token>")
def verify_email(token):

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    print("Token received:", token)

    cursor.execute(
        """
        SELECT
            email,
            expires_at,
            used
        FROM email_verification_tokens
        WHERE token=?
        """,
        (token,)
    )

    token_data = cursor.fetchone()

    print("Database result:", token_data)

    if not token_data:

        print("Token NOT found")

        conn.close()

        return render_template(
            "invalid_verification_link.html"
        )

    email = token_data[0]
    expires_at = token_data[1]
    used = token_data[2]

    print("Email:", email)
    print("Expires:", expires_at)
    print("Used:", used)

    expires_at = datetime.fromisoformat(expires_at)

    if datetime.utcnow() > expires_at:

        print("Token expired")

        conn.close()

        return render_template(
            "expired_verification_link.html"
        )

    if used == 1:

        print("Token already used")

        conn.close()

        return render_template(
            "invalid_verification_link.html"
        )

    print("Updating user...")

    cursor.execute(
        """
        UPDATE users
        SET email_verified=1
        WHERE email=?
        """,
        (email,)
    )

    cursor.execute(
        """
        UPDATE email_verification_tokens
        SET used=1
        WHERE token=?
        """,
        (token,)
    )

    conn.commit()

    print("Verification completed successfully")

    conn.close()

    return render_template(
        "email_verified.html"
    )




# ✅ Dashboard with Analytics
@app.route("/dashboard")
def dashboard():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Total exams taken
    cursor.execute(
        """
        SELECT COUNT(*)
        FROM results
        WHERE username=?
        """,
        (session["user"],)
    )

    total_exams = cursor.fetchone()[0]

    # Average percentage
    cursor.execute(
        """
        SELECT AVG(percentage)
        FROM results
        WHERE username=?
        """,
        (session["user"],)
    )

    avg_result = cursor.fetchone()[0]
    average_percentage = round(avg_result, 1) if avg_result else 0

    # Highest percentage
    cursor.execute(
        """
        SELECT MAX(percentage)
        FROM results
        WHERE username=?
        """,
        (session["user"],)
    )

    best_result = cursor.fetchone()[0]
    highest_percentage = best_result if best_result else 0

    # Recent exams
    cursor.execute(
        """
        SELECT
            id,
            exam_type,
            exam_name,
            percentage,
            date_taken
        FROM results
        WHERE username=?
        ORDER BY id DESC
        LIMIT 5
        """,
        (session["user"],)
    )

    recent_exams = cursor.fetchall()
    
    # Total Bookmarks
    cursor.execute(
        """
        SELECT COUNT(*)
        FROM bookmarks
        WHERE username=?
        """,
        (session["user"],)
    )

    total_bookmarks = cursor.fetchone()[0]
    
    # -----------------------------------
    # Subscription Information
    # -----------------------------------

    from datetime import datetime

    cursor.execute(
        """
        SELECT
            plan_name,
            start_date,
            end_date,
            is_active
        FROM subscriptions
        WHERE username=?
        """,
        (session["user"],)
    )

    subscription = cursor.fetchone()

    subscription_plan = "Free"
    subscription_status = "Inactive"
    expiry_date = None
    days_remaining = 0
    show_expiry_warning = False

    if subscription:

        subscription_plan = subscription[0]
        expiry_date = subscription[2]
        is_active = subscription[3]

        if is_active and expiry_date:

            expiry = datetime.strptime(
                expiry_date,
                "%Y-%m-%d %H:%M:%S"
            )

            now = datetime.now()

            if expiry > now:

                subscription_status = "Active"

                days_remaining = (expiry - now).days

                # Show warning if 7 days or less remain
                if days_remaining <= 7:
                    show_expiry_warning = True

            else:

                subscription_status = "Expired"

                cursor.execute(
                    """
                    UPDATE subscriptions
                    SET is_active=0
                    WHERE username=?
                    """,
                    (session["user"],)
                )

                conn.commit()

    conn.close()

    return render_template(
        "dashboard.html",
        user=session.get("name", session.get("user")),
        total_exams=total_exams,
        average_percentage=average_percentage,
        highest_percentage=highest_percentage,
        total_bookmarks=total_bookmarks,
        recent_exams=recent_exams,
        subscription_plan=subscription_plan,
        subscription_status=subscription_status,
        expiry_date=expiry_date,
        days_remaining=days_remaining,
        show_expiry_warning=show_expiry_warning
    )
    
    
    
# ✅ Practice Mode Home
@app.route("/practice")
def practice():

    if "user" not in session:
        return redirect("/login")

    return render_template("practice.html")



# ✅ WAEC Practice
@app.route("/practice_waec")
def practice_waec():

    if "user" not in session:
        return redirect("/login")

    if not can_take_exam(session["user"]):
        return redirect("/subscribe")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT DISTINCT subject
        FROM questions
        WHERE exam_type='WAEC'
        ORDER BY subject
        """
    )

    subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "practice_waec.html",
        subjects=subjects
    )


# ✅ Practice Questions
@app.route(
    "/practice_questions/<exam_type>/<subject>",
    methods=["GET", "POST"]
)
def practice_questions(exam_type, subject):

    if "user" not in session:
        return redirect("/login")

    if not can_take_exam(session["user"]):
        return redirect("/subscribe")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    feedback = False
    selected_answer = None
    correct_answer = None
    explanation = None
    is_correct = None

    # ==========================
    # GET → Load random question
    # ==========================

    if request.method == "GET":

        cursor.execute(
            """
            SELECT
                id,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation
            FROM questions
            WHERE exam_type=?
            AND subject=?
            ORDER BY RANDOM()
            LIMIT 1
            """,
            (
                exam_type,
                subject
            )
        )

        question = cursor.fetchone()

    # ==========================
    # POST → Reload SAME question
    # ==========================

    else:

        question_id = request.form["question_id"]

        cursor.execute(
            """
            SELECT
                id,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation
            FROM questions
            WHERE id=?
            """,
            (
                question_id,
            )
        )

        question = cursor.fetchone()

        selected_answer = request.form["answer"]

        correct_answer = question[6]

        explanation = question[7]

        is_correct = (
            selected_answer == correct_answer
        )

        feedback = True

    conn.close()

    if not question:

        return "No practice questions found."

    return render_template(
        "practice_question.html",
        exam_type=exam_type,
        subject=subject,
        question=question,
        feedback=feedback,
        selected_answer=selected_answer,
        correct_answer=correct_answer,
        explanation=explanation,
        is_correct=is_correct
    )



# ✅ Student Leaderboard
@app.route("/leaderboard")
def leaderboard():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            username,
            exam_type,
            exam_name,
            percentage,
            date_taken
        FROM results
        ORDER BY percentage DESC, date_taken DESC
        LIMIT 10
        """
    )

    leaders = cursor.fetchall()

    conn.close()

    return render_template(
        "leaderboard.html",
        leaders=leaders
    )
    
# ✅ Exam type page
@app.route("/exam_types")
def exam_types():
    return render_template("exam_types.html")


# ✅ Dynamic POST-UTME Universities
@app.route("/post_utme")
def post_utme():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT university_name, exam_mode
        FROM post_utme_universities
        ORDER BY university_name
    """)

    universities = cursor.fetchall()

    conn.close()

    return render_template(
        "post_utme.html",
        universities=universities
    )


# ✅ POST-UTME Mode Handler
@app.route("/post_utme_handler")
def post_utme_handler():

    if "user" not in session:
        return redirect("/login")

    university = request.args.get("university")

    if not university:
        return redirect("/post_utme")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT exam_mode
        FROM post_utme_universities
        WHERE university_name=?
        """,
        (university,)
    )

    result = cursor.fetchone()

    conn.close()

    if not result:
        return "University not found."

    exam_mode = result[0]

    if exam_mode == "CBT":
        return redirect(
            f"/post_utme_courses?university={university}"
        )

    elif exam_mode == "APTITUDE":
        return redirect(
            f"/post_utme_subjects?university={university}"
        )

    elif exam_mode == "CURRENT_AFFAIRS":
        return redirect(
            f"/post_utme_subjects?university={university}"
        )

    elif exam_mode == "SCREENING":

        return render_template(
            "screening_notice.html",
            university=university
        )

    return "Unsupported exam mode."   


# ✅ POST-UTME Course Selection
@app.route("/post_utme_courses")
def post_utme_courses():

    if "user" not in session:
        return redirect("/login")

    university = request.args.get("university")

    if not university:
        return redirect("/post_utme")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT course_name
        FROM post_utme_courses
        WHERE university_name=?
        ORDER BY course_name
        """,
        (university,)
    )

    courses = cursor.fetchall()

    conn.close()

    return render_template(
        "post_utme_courses.html",
        university=university,
        courses=courses
    ) 


# ✅ POST-UTME Subject Combination
@app.route("/post_utme_subjects")
def post_utme_subjects():

    if "user" not in session:
        return redirect("/login")

    university = request.args.get("university")
    course = request.args.get("course")

    if not university:
        return redirect("/post_utme")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get university exam mode
    cursor.execute(
        """
        SELECT exam_mode
        FROM post_utme_universities
        WHERE university_name=?
        """,
        (university,)
    )

    mode_result = cursor.fetchone()

    if not mode_result:
        conn.close()
        return "University not found."

    exam_mode = mode_result[0]

    # CBT universities
    if exam_mode == "CBT":

        if not course:
            conn.close()
            return redirect(
                f"/post_utme_courses?university={university}"
            )

        cursor.execute(
            """
            SELECT subject_name
            FROM post_utme_course_subjects
            WHERE course_name=?
            """,
            (course,)
        )

        subjects = cursor.fetchall()

    else:

        cursor.execute(
            """
            SELECT subject_name
            FROM post_utme_subjects
            WHERE university_name=?
            """,
            (university,)
        )

        subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "post_utme_subjects.html",
        university=university,
        course=course,
        exam_mode=exam_mode,
        subjects=subjects
    )


# ✅ Add POST-UTME Questions
@app.route("/add_post_utme_question", methods=["GET", "POST"])
def add_post_utme_question():

    if request.method == "POST":

        university = request.form["university"]
        subject = request.form["subject"]

        question = request.form["question"]

        option_a = request.form["option_a"]
        option_b = request.form["option_b"]
        option_c = request.form["option_c"]
        option_d = request.form["option_d"]

        correct_answer = request.form["correct_answer"]

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            INSERT INTO post_utme_questions (
                university_name,
                subject,
                question,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                university,
                subject,
                question,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer
            )
        )

        conn.commit()
        conn.close()

        return "POST-UTME Question Added Successfully."

    # Load universities for dropdown
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT university_name
        FROM post_utme_universities
        ORDER BY university_name
        """
    )

    universities = cursor.fetchall()

    conn.close()

    return render_template(
        "add_post_utme_question.html",
        universities=universities
    )


# ✅ Get subjects for selected university
@app.route("/get_post_utme_subjects")
def get_post_utme_subjects():

    university = request.args.get("university")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT subject_name
        FROM post_utme_subjects
        WHERE university_name=?
        ORDER BY subject_name
        """,
        (university,)
    )

    subjects = cursor.fetchall()

    conn.close()

    return {
        "subjects": [s[0] for s in subjects]
    }
    
    
# ✅ Start POST-UTME Exam
@app.route("/start_post_utme")
def start_post_utme():

    if "user" not in session:
        return redirect("/login")

    # Check subscription
    if not can_take_exam(session["user"]):
        return redirect("/subscribe")

    university = request.args.get("university")
    course = request.args.get("course")

    if not university:
        return redirect("/post_utme")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT exam_mode
        FROM post_utme_universities
        WHERE university_name=?
        """,
        (university,)
    )

    row = cursor.fetchone()

    if not row:
        conn.close()
        return "University not found."

    exam_mode = row[0]

    # CBT universities
    if exam_mode == "CBT":

        cursor.execute(
            """
            SELECT subject_name
            FROM post_utme_course_subjects
            WHERE course_name=?
            """,
            (course,)
        )

    else:

        cursor.execute(
            """
            SELECT subject_name
            FROM post_utme_subjects
            WHERE university_name=?
            """,
            (university,)
        )

    subjects = [x[0] for x in cursor.fetchall()]

    conn.close()

    if not subjects:
        return "No subjects found."

    session["exam_type"] = "POST-UTME"

    session["university"] = university
    session["course"] = course

    session["subjects"] = subjects
    session["subject_index"] = 0
    session["subject"] = subjects[0]

    session["q_index"] = 0
    session["score"] = 0
    session["total_answered"] = 0

    session["subject_scores"] = {}
    session["subject_totals"] = {}

    session["start_time"] = time.time()
    session["exam_duration"] = 40 * 60

    return redirect(
        f"/question/POST-UTME/{subjects[0]}"
    )


# ✅ Manage POST-UTME Universities
@app.route("/manage_post_utme_universities")
def manage_post_utme_universities():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            university_name,
            exam_mode,
            duration,
            total_questions,
            pass_mark
        FROM post_utme_universities
        ORDER BY university_name
    """)

    universities = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_post_utme_universities.html",
        universities=universities
    )


# ✅ Edit POST-UTME University
@app.route(
    "/edit_post_utme_university/<int:university_id>",
    methods=["GET", "POST"]
)
def edit_post_utme_university(university_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        university_name = request.form["university_name"]
        exam_mode = request.form["exam_mode"]

        duration = request.form["duration"]
        total_questions = request.form["total_questions"]
        pass_mark = request.form["pass_mark"]

        cursor.execute(
            """
            UPDATE post_utme_universities
            SET
                university_name=?,
                exam_mode=?,
                duration=?,
                total_questions=?,
                pass_mark=?
            WHERE id=?
            """,
            (
                university_name,
                exam_mode,
                duration,
                total_questions,
                pass_mark,
                university_id
            )
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_post_utme_universities"
        )

    cursor.execute(
        """
        SELECT *
        FROM post_utme_universities
        WHERE id=?
        """,
        (university_id,)
    )

    university = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_post_utme_university.html",
        university=university
    )


# ✅ Manage POST-UTME Courses
@app.route("/manage_post_utme_courses")
def manage_post_utme_courses():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get(
        "search",
        ""
    ).strip()

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if search:

        cursor.execute(
            """
            SELECT *
            FROM post_utme_courses
            WHERE
                university_name LIKE ?
                OR course_name LIKE ?
            ORDER BY university_name
            """,
            (
                f"%{search}%",
                f"%{search}%"
            )
        )

    else:

        cursor.execute(
            """
            SELECT *
            FROM post_utme_courses
            ORDER BY university_name
            """
        )

    courses = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_post_utme_courses.html",
        courses=courses,
        search=search
    )


# ✅ Edit POST-UTME Course
@app.route(
    "/edit_post_utme_course/<int:course_id>",
    methods=["GET", "POST"]
)
def edit_post_utme_course(course_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        university_name = request.form["university_name"]
        course_name = request.form["course_name"]

        cursor.execute(
            """
            UPDATE post_utme_courses
            SET
                university_name=?,
                course_name=?
            WHERE id=?
            """,
            (
                university_name,
                course_name,
                course_id
            )
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_post_utme_courses"
        )

    # Current course
    cursor.execute(
        """
        SELECT *
        FROM post_utme_courses
        WHERE id=?
        """,
        (course_id,)
    )

    course = cursor.fetchone()

    # Universities dropdown
    cursor.execute(
        """
        SELECT university_name
        FROM post_utme_universities
        ORDER BY university_name
        """
    )

    universities = cursor.fetchall()

    conn.close()

    return render_template(
        "edit_post_utme_course.html",
        course=course,
        universities=universities
    )


# ✅ Add POST-UTME Course
@app.route(
    "/add_post_utme_course",
    methods=["GET", "POST"]
)
def add_post_utme_course():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        university_name = request.form[
            "university_name"
        ]

        course_name = request.form[
            "course_name"
        ]

        # Prevent duplicate course
        cursor.execute(
            """
            SELECT id
            FROM post_utme_courses
            WHERE
                university_name=?
                AND course_name=?
            """,
            (
                university_name,
                course_name
            )
        )

        existing = cursor.fetchone()

        if existing:

            conn.close()

            return """
            Course already exists.
            """

        cursor.execute(
            """
            INSERT INTO post_utme_courses
            (
                university_name,
                course_name
            )
            VALUES (?, ?)
            """,
            (
                university_name,
                course_name
            )
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_post_utme_courses"
        )

    cursor.execute(
        """
        SELECT university_name
        FROM post_utme_universities
        ORDER BY university_name
        """
    )

    universities = cursor.fetchall()

    conn.close()

    return render_template(
        "add_post_utme_course.html",
        universities=universities
    )



# ==========================================
# SUBJECT MANAGEMENT
# ==========================================

@app.route("/manage_subjects")
def manage_subjects():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get("search", "").strip()

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    sql = """
        SELECT *
        FROM subjects
    """

    params = []

    if search:

        sql += """
            WHERE subject_name LIKE ?
            OR exam_type LIKE ?
            OR subject_code LIKE ?
        """

        keyword = f"%{search}%"

        params = [
            keyword,
            keyword,
            keyword
        ]

    sql += """
        ORDER BY exam_type,
                 subject_name
    """

    cursor.execute(sql, params)

    subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_subjects.html",
        subjects=subjects,
        search=search
    )



# ==========================================
# ADD SUBJECT
# ==========================================

@app.route("/add_subject", methods=["GET", "POST"])
def add_subject():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        exam_type = request.form["exam_type"].strip()
        subject_name = request.form["subject_name"].strip()
        subject_code = request.form["subject_code"].strip().upper()
        status = request.form["status"].strip()

        # Validation
        if not exam_type or not subject_name:
            conn.close()
            return "Exam Type and Subject Name are required."

        # Prevent duplicate subjects
        cursor.execute("""
            SELECT id
            FROM subjects
            WHERE exam_type=?
            AND subject_name=?
        """, (exam_type, subject_name))

        existing = cursor.fetchone()

        if existing:

            conn.close()

            flash(
                "Subject already exists.",
                "warning"
            )

            return redirect("/manage_subjects")

        cursor.execute("""
            INSERT INTO subjects
            (
                exam_type,
                subject_name,
                subject_code,
                status
            )
            VALUES (?, ?, ?, ?)
        """, (
            exam_type,
            subject_name,
            subject_code,
            status
        ))

        conn.commit()
        conn.close()

        flash(
            "Subject added successfully.",
            "success"
        )

        return redirect("/manage_subjects")

    conn.close()

    return render_template("add_subject.html")



# ==========================================
# EDIT SUBJECT
# ==========================================

@app.route("/edit_subject/<int:subject_id>", methods=["GET", "POST"])
def edit_subject(subject_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    if request.method == "POST":

        exam_type = request.form["exam_type"].strip()
        subject_name = request.form["subject_name"].strip()
        subject_code = request.form["subject_code"].strip().upper()
        status = request.form["status"].strip()

        cursor.execute("""
            UPDATE subjects
            SET
                exam_type=?,
                subject_name=?,
                subject_code=?,
                status=?
            WHERE id=?
        """, (
            exam_type,
            subject_name,
            subject_code,
            status,
            subject_id
        ))

        conn.commit()
        conn.close()
        
        flash(
            "Subject updated successfully.",
            "success"
        )

        return redirect("/manage_subjects")

    cursor.execute("""
        SELECT *
        FROM subjects
        WHERE id=?
    """, (subject_id,))

    subject = cursor.fetchone()

    conn.close()

    if not subject:
        return "Subject not found."

    return render_template(
        "edit_subject.html",
        subject=subject
    )



# ==========================================
# ACTIVATE / DEACTIVATE SUBJECT
# ==========================================

@app.route("/toggle_subject/<int:subject_id>")
def toggle_subject(subject_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT status
        FROM subjects
        WHERE id=?
    """, (subject_id,))

    row = cursor.fetchone()

    if not row:

        conn.close()

        flash(
            "Subject not found.",
            "danger"
        )

        return redirect("/manage_subjects")

    if row[0] == "Active":

        new_status = "Inactive"
        flash_message = "✅ Subject deactivated successfully."

    else:

        new_status = "Active"
        flash_message = "✅ Subject activated successfully."

    cursor.execute("""
        UPDATE subjects
        SET status=?
        WHERE id=?
    """, (
        new_status,
        subject_id
    ))

    conn.commit()
    conn.close()

    flash(
        flash_message,
        "success"
    )

    return redirect("/manage_subjects")



# ==========================================
# DELETE SUBJECT
# ==========================================

@app.route("/delete_subject/<int:subject_id>", methods=["POST"])
def delete_subject(subject_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # ------------------------------------------
    # Check that subject exists
    # ------------------------------------------

    cursor.execute(
        """
        SELECT subject_name, exam_type
        FROM subjects
        WHERE id=?
        """,
        (subject_id,)
    )

    subject = cursor.fetchone()

    if not subject:
        conn.close()
        return "Subject not found.", 404

    subject_name = subject[0]
    exam_type = subject[1]

    # ------------------------------------------
    # Check for topics
    # ------------------------------------------

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM topics
        WHERE subject_id=?
        """,
        (subject_id,)
    )

    topic_count = cursor.fetchone()[0]

    if topic_count > 0:

        conn.close()

        return (
            f"Cannot delete {exam_type} - {subject_name}. "
            f"It has {topic_count} topic(s) attached. "
            f"Remove the topics first."
        )

    # ------------------------------------------
    # Check for questions
    # ------------------------------------------

    cursor.execute(
        """
        SELECT COUNT(*)
        FROM questions_v2
        WHERE subject_id=?
        """,
        (subject_id,)
    )

    question_count = cursor.fetchone()[0]

    if question_count > 0:

        conn.close()

        return (
            f"Cannot delete {exam_type} - {subject_name}. "
            f"It has {question_count} question(s) attached."
        )

    # ------------------------------------------
    # Delete subject
    # ------------------------------------------

    cursor.execute(
        """
        DELETE FROM subjects
        WHERE id=?
        """,
        (subject_id,)
    )

    conn.commit()
    conn.close()

    return redirect("/manage_subjects")



# ==========================================
# MANAGE TOPICS
# ==========================================

@app.route("/manage_topics")
def manage_topics():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get(
        "search",
        ""
    ).strip()

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    query = """
        SELECT
            topics.id,
            topics.topic_name,
            topics.status,
            topics.created_at,
            subjects.subject_name,
            subjects.exam_type
        FROM topics
        INNER JOIN subjects
            ON topics.subject_id = subjects.id
    """

    params = []

    if search:

        query += """
            WHERE
                topics.topic_name LIKE ?
                OR subjects.subject_name LIKE ?
                OR subjects.exam_type LIKE ?
        """

        keyword = f"%{search}%"

        params.extend([
            keyword,
            keyword,
            keyword
        ])

    query += """
        ORDER BY
            subjects.exam_type,
            subjects.subject_name,
            topics.topic_name
    """

    cursor.execute(query, params)

    topics = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_topics.html",
        topics=topics,
        search=search
    )



# ==========================================
# ADD TOPIC
# ==========================================

@app.route("/add_topic", methods=["GET", "POST"])
def add_topic():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    if request.method == "POST":

        subject_id = request.form["subject_id"]
        topic_name = request.form["topic_name"].strip()

        # Prevent duplicate topics under the same subject
        cursor.execute(
            """
            SELECT id
            FROM topics
            WHERE subject_id=?
            AND topic_name=?
            """,
            (
                subject_id,
                topic_name
            )
        )

        existing = cursor.fetchone()

        if existing:

            conn.close()

            flash(
                "Topic already exists.",
                "warning"
            )

            return redirect("/add_topic")

        cursor.execute(
            """
            INSERT INTO topics
            (
                subject_id,
                topic_name
            )
            VALUES (?, ?)
            """,
            (
                subject_id,
                topic_name
            )
        )

        conn.commit()
        conn.close()

        flash(
            "Topic added successfully.",
            "success"
        )

        return redirect("/manage_topics")

    cursor.execute("""
        SELECT
            id,
            exam_type,
            subject_name
        FROM subjects
        WHERE status='Active'
        ORDER BY
            exam_type,
            subject_name
    """)

    subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "add_topic.html",
        subjects=subjects
    )



# ==========================================
# EDIT TOPIC
# ==========================================

@app.route("/edit_topic/<int:topic_id>", methods=["GET", "POST"])
def edit_topic(topic_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    if request.method == "POST":

        subject_id = request.form["subject_id"]
        topic_name = request.form["topic_name"].strip()
        
        # ------------------------------------
        # Check for duplicate topic
        # ------------------------------------

        cursor.execute("""
            SELECT id
            FROM topics
            WHERE
                subject_id=?
                AND topic_name=?
                AND id<>?
        """, (
            subject_id,
            topic_name,
            topic_id
        ))

        existing = cursor.fetchone()

        if existing:

            conn.close()

            flash(
                "⚠️ Topic already exists for this subject.",
                "warning"
            )

            return redirect(f"/edit_topic/{topic_id}")
        

        cursor.execute("""
            UPDATE topics
            SET
                subject_id=?,
                topic_name=?
            WHERE id=?
        """, (
            subject_id,
            topic_name,
            topic_id
        ))

        conn.commit()
        conn.close()

        flash(
            "Topic updated successfully.",
            "success"
        )

        return redirect("/manage_topics")

    cursor.execute("""
        SELECT *
        FROM topics
        WHERE id=?
    """, (topic_id,))

    topic = cursor.fetchone()

    if not topic:

        conn.close()

        return "Topic not found."

    cursor.execute("""
        SELECT
            id,
            exam_type,
            subject_name
        FROM subjects
        WHERE status='Active'
        ORDER BY
            exam_type,
            subject_name
    """)

    subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "edit_topic.html",
        topic=topic,
        subjects=subjects
    )


# ==========================================
# ACTIVATE / DEACTIVATE TOPIC
# ==========================================

@app.route("/toggle_topic/<int:topic_id>")
def toggle_topic(topic_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT status
        FROM topics
        WHERE id=?
    """, (topic_id,))

    row = cursor.fetchone()

    if not row:

        conn.close()

        flash(
            "Topic not found.",
            "danger"
        )

        return redirect("/manage_topics")

    if row[0] == "Active":

        new_status = "Inactive"
        flash_message = "✅ Topic deactivated successfully."

    else:

        new_status = "Active"
        flash_message = "✅ Topic activated successfully."

    cursor.execute("""
        UPDATE topics
        SET status=?
        WHERE id=?
    """, (
        new_status,
        topic_id
    ))

    conn.commit()
    conn.close()

    flash(
        flash_message,
        "success"
    )

    return redirect("/manage_topics")



# ✅ Manage Subject Combinations
@app.route("/manage_subject_combinations")
def manage_subject_combinations():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get(
        "search",
        ""
    ).strip()

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if search:

        cursor.execute(
            """
            SELECT *
            FROM post_utme_course_subjects
            WHERE
                course_name LIKE ?
                OR subject_name LIKE ?
            ORDER BY course_name
            """,
            (
                f"%{search}%",
                f"%{search}%"
            )
        )

    else:

        cursor.execute(
            """
            SELECT *
            FROM post_utme_course_subjects
            ORDER BY course_name
            """
        )

    combinations = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_subject_combinations.html",
        combinations=combinations,
        search=search
    )


# ✅ Add Subject Combination
@app.route(
    "/add_subject_combination",
    methods=["GET", "POST"]
)
def add_subject_combination():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        course_name = request.form["course_name"]
        subject_name = request.form["subject_name"]

        # Prevent duplicates
        cursor.execute(
            """
            SELECT id
            FROM post_utme_course_subjects
            WHERE
                course_name=?
                AND subject_name=?
            """,
            (
                course_name,
                subject_name
            )
        )

        existing = cursor.fetchone()

        if existing:

            conn.close()

            return "Subject combination already exists."

        cursor.execute(
            """
            INSERT INTO post_utme_course_subjects
            (
                course_name,
                subject_name
            )
            VALUES (?, ?)
            """,
            (
                course_name,
                subject_name
            )
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_subject_combinations"
        )

    cursor.execute(
        """
        SELECT DISTINCT course_name
        FROM post_utme_courses
        ORDER BY course_name
        """
    )

    courses = cursor.fetchall()

    conn.close()

    return render_template(
        "add_subject_combination.html",
        courses=courses
    )


# ✅ Edit Subject Combination
@app.route(
    "/edit_subject_combination/<int:combination_id>",
    methods=["GET", "POST"]
)
def edit_subject_combination(combination_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        course_name = request.form["course_name"]
        subject_name = request.form["subject_name"]

        cursor.execute(
            """
            UPDATE post_utme_course_subjects
            SET
                course_name=?,
                subject_name=?
            WHERE id=?
            """,
            (
                course_name,
                subject_name,
                combination_id
            )
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_subject_combinations"
        )

    # Current record
    cursor.execute(
        """
        SELECT *
        FROM post_utme_course_subjects
        WHERE id=?
        """,
        (combination_id,)
    )

    combination = cursor.fetchone()

    # Course dropdown
    cursor.execute(
        """
        SELECT DISTINCT course_name
        FROM post_utme_courses
        ORDER BY course_name
        """
    )

    courses = cursor.fetchall()

    conn.close()

    return render_template(
        "edit_subject_combination.html",
        combination=combination,
        courses=courses
    )


# ✅ JAMB course selection
@app.route("/jamb_courses")
def jamb_courses():

    return render_template(
        "jamb_courses.html",
        courses=JAMB_COURSES
    )
    

# ✅ WAEC subject selection
@app.route("/waec_subjects")
def waec_subjects():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT DISTINCT subject
        FROM questions
        WHERE exam_type='WAEC'
    """)

    subjects_data = cursor.fetchall()

    conn.close()

    subjects = []

    for sub in subjects_data:
        subjects.append(sub[0])

    return render_template(
        "waec_subjects.html",
        subjects=subjects
    )
    
# ✅ Start WAEC subject CBT
@app.route("/start_waec/<subject>")
def start_waec(subject):

    # Check if the student already has an unfinished exam
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT exam_type, subject
        FROM exam_progress
        WHERE username=?
        AND status='IN_PROGRESS'
        """,
        (
            session.get("user"),
        )
    )

    unfinished = cursor.fetchone()

    conn.close()

    if unfinished:

        return render_template(
            "resume_exam.html",
            exam_type=unfinished[0],
            subject=unfinished[1]
        )

   # Check subscription
    if not can_take_exam(session["user"]):
        return redirect("/subscribe")
   
    # Start a new WAEC exam
    session["exam_type"] = "WAEC"
        
    session["subject"] = subject

    session["q_index"] = 0
    session["score"] = 0
    session["total_answered"] = 0

    session["subject_scores"] = {}
    session["subject_totals"] = {}

    session["start_time"] = __import__("time").time()
    session["exam_duration"] = 40 * 60

    # Save exam progress to database
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Check if this exact exam already exists
    cursor.execute(
        """
        SELECT id
        FROM exam_progress
        WHERE username=?
        AND exam_type=?
        AND subject=?
        AND status='IN_PROGRESS'
        """,
        (
            session.get("user"),
            "WAEC",
            subject
        )
    )

    existing = cursor.fetchone()

    if existing:

        cursor.execute(
            """
            UPDATE exam_progress
            SET
                exam_name=?,
                subject=?,
                current_subject_index=?,
                q_index=?,
                score=?,
                total_answered=?,
                start_time=?,
                exam_duration=?,
                updated_at=CURRENT_TIMESTAMP
            WHERE id=?
            """,
            (
                subject,
                subject,
                0,
                0,
                0,
                0,
                session["start_time"],
                session["exam_duration"],
                existing[0]
            )
        )

    else:

        cursor.execute(
            """
            INSERT INTO exam_progress
            (
                username,
                exam_type,
                exam_name,
                subject,
                current_subject_index,
                q_index,
                score,
                total_answered,
                start_time,
                exam_duration,
                status
            )
            VALUES
            (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                session.get("user"),
                "WAEC",
                subject,
                subject,
                0,
                0,
                0,
                0,
                session["start_time"],
                session["exam_duration"],
                "IN_PROGRESS"
            )
        )

    conn.commit()
    conn.close()

    return redirect(
        f"/question/WAEC/{subject}"
    )
    
    
# ✅ Subscription Page
@app.route("/subscribe")
def subscribe():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            plan_name,
            price,
            duration_days,
            description
        FROM subscription_plans
        WHERE is_active = 1
        ORDER BY duration_days
        """
    )

    plans = cursor.fetchall()

    conn.close()

    return render_template(
        "subscribe.html",
        plans=plans
    )
    
    
# ✅ Subscription Summary
@app.route("/subscribe_plan/<int:plan_id>")
def subscribe_plan(plan_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            plan_name,
            price,
            duration_days,
            description
        FROM subscription_plans
        WHERE id=?
        AND is_active=1
        """,
        (plan_id,)
    )

    plan = cursor.fetchone()

    conn.close()

    if not plan:
        return "Subscription plan not found."

    return render_template(
        "subscription_summary.html",
        plan=plan
    )
    
    

# ✅ Initialize Paystack Payment
@app.route("/initialize_payment/<int:plan_id>")
def initialize_payment(plan_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get selected plan
    cursor.execute(
        """
        SELECT
            id,
            plan_name,
            price,
            duration_days
        FROM subscription_plans
        WHERE id=?
        AND is_active=1
        """,
        (plan_id,)
    )

    plan = cursor.fetchone()

    if not plan:

        conn.close()
        return "Subscription plan not found."

    transaction_reference = (
        f"CBT-{int(time.time())}-{secrets.token_hex(8)}"
    )

    # Save pending payment
    cursor.execute(
        """
        INSERT INTO payments
        (
            username,
            plan_id,
            plan_name,
            amount,
            duration_days,
            transaction_reference,
            payment_status
        )
        VALUES
        (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            session["user"],
            plan[0],
            plan[1],
            plan[2],
            plan[3],
            transaction_reference,
            "PENDING"
        )
    )

    conn.commit()
    conn.close()

    headers = {
        "Authorization": f"Bearer {PAYSTACK_SECRET_KEY}",
        "Content-Type": "application/json"
    }

    data = {
        "email": session["user"],
        "amount": int(plan[2] * 100),   # Kobo
        "reference": transaction_reference,
        "callback_url": f"{APP_URL}/payment_callback"
    }

    try:
        response = requests.post(
            "https://api.paystack.co/transaction/initialize",
            json=data,
            headers=headers,
            timeout=30
        )

    except requests.exceptions.RequestException:

        return "Unable to connect to Paystack. Please try again."

    try:
        result = response.json()

    except ValueError:

        return "Invalid response received from Paystack."
        
    if response.status_code != 200:
        return f"Paystack Error: {result}"

    if result["status"]:

        return redirect(
            result["data"]["authorization_url"]
        )

    return result["message"]



# ✅ Paystack Payment Callback
@app.route("/payment_callback")
def payment_callback():

    reference = request.args.get("reference")

    if not reference:
        return "Payment reference not found."

    headers = {
        "Authorization": f"Bearer {PAYSTACK_SECRET_KEY}"
    }

    response = requests.get(
        f"https://api.paystack.co/transaction/verify/{reference}",
        headers=headers,
        timeout=30
    )

    print("Status Code:", response.status_code)
    print("Response Text:")
    print(response.text)

    try:
        result = response.json()

    except ValueError:
        return "Invalid response received from Paystack."

    if response.status_code != 200:
        return f"Verification failed: {result}"

    if not result["status"]:
        return result["message"]

    payment = result["data"]

    # Prevent processing the same payment twice
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT payment_status
        FROM payments
        WHERE transaction_reference=?
        """,
        (reference,)
    )

    existing_payment = cursor.fetchone()

    if existing_payment and existing_payment[0] == "SUCCESS":
        conn.close()
        return redirect("/dashboard")

    conn.close()

    if payment["status"] != "success":
        return "Payment was not successful."

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        UPDATE payments
        SET
            payment_status=?,
            paystack_reference=?,
            gateway_response=?,
            payment_method=?,
            currency=?,
            verified_at=CURRENT_TIMESTAMP,
            paid_at=CURRENT_TIMESTAMP
        WHERE transaction_reference=?
        """,
        (
            "SUCCESS",
            payment["reference"],
            payment["gateway_response"],
            payment["channel"],
            payment["currency"],
            reference
        )
    )

    conn.commit()

    # ------------------------------------
    # Get payment details
    # ------------------------------------

    cursor.execute(
        """
        SELECT
            username,
            plan_id,
            plan_name,
            duration_days,
            amount,
            currency,
            transaction_reference
        FROM payments
        WHERE transaction_reference=?
        """,
        (reference,)
    )

    payment_record = cursor.fetchone()

    if not payment_record:
        conn.close()
        return "Payment record not found."

    username = payment_record[0]
    plan_id = payment_record[1]
    plan_name = payment_record[2]
    duration_days = payment_record[3]
    amount_paid = payment_record[4]
    currency = payment_record[5]
    payment_reference = payment_record[6]

    from datetime import datetime, timedelta

    now = datetime.now()

    cursor.execute(
        """
        SELECT id
        FROM subscriptions
        WHERE username=?
        """,
        (username,)
    )

    existing = cursor.fetchone()

    if existing:

        cursor.execute(
            """
            SELECT end_date
            FROM subscriptions
            WHERE username=?
            """,
            (username,)
        )

        current = cursor.fetchone()

        if current and current[0]:

            current_end = datetime.strptime(
                current[0],
                "%Y-%m-%d %H:%M:%S"
            )

            if current_end > now:
                new_end = current_end + timedelta(days=duration_days)
            else:
                new_end = now + timedelta(days=duration_days)

        else:

            new_end = now + timedelta(days=duration_days)

        cursor.execute(
            """
            UPDATE subscriptions
            SET
                plan_id=?,
                plan_name=?,
                start_date=?,
                end_date=?,
                payment_reference=?,
                payment_status='SUCCESS',
                is_active=1,
                amount_paid=?,
                currency=?
            WHERE username=?
            """,
            (
                plan_id,
                plan_name,
                now.strftime("%Y-%m-%d %H:%M:%S"),
                new_end.strftime("%Y-%m-%d %H:%M:%S"),
                payment_reference,
                amount_paid,
                currency,
                username
            )
        )

    else:

        new_end = now + timedelta(days=duration_days)

        cursor.execute(
            """
            INSERT INTO subscriptions
            (
                username,
                plan_id,
                plan_name,
                start_date,
                end_date,
                payment_reference,
                payment_status,
                is_active,
                amount_paid,
                currency
            )
            VALUES
            (?, ?, ?, ?, ?, ?, 'SUCCESS', 1, ?, ?)
            """,
            (
                username,
                plan_id,
                plan_name,
                now.strftime("%Y-%m-%d %H:%M:%S"),
                new_end.strftime("%Y-%m-%d %H:%M:%S"),
                payment_reference,
                amount_paid,
                currency
            )
        )

    # ------------------------------------
    # Get Student Name
    # ------------------------------------

    cursor.execute(
        """
        SELECT name
        FROM users
        WHERE email=?
        """,
        (username,)
    )

    student = cursor.fetchone()

    if student:
        student_name = student[0]
    else:
        student_name = "Student"

    conn.commit()
    conn.close()

    # =====================================
    # Send Payment Confirmation Email
    # =====================================

    email_body = f"""
Hello {student_name},

Thank you for subscribing to PrepNova CBT.

Your payment has been received successfully.

----------------------------------------

Subscription Plan:
{plan_name}

Amount Paid:
₦{amount_paid:,.0f}

Payment Reference:
{payment_reference}

Subscription Starts:
{now.strftime("%d %B %Y")}

Subscription Expires:
{new_end.strftime("%d %B %Y")}

Status:
ACTIVE

----------------------------------------

You can now continue practicing for your examinations.

Thank you for choosing PrepNova CBT.

Regards,

PrepNova CBT Team
"""

    send_email(
        subject="PrepNova CBT Payment Confirmation",
        recipient=username,
        body=email_body
    )

    return redirect("/dashboard")


@app.route("/check_payments_table")
def check_payments_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("PRAGMA table_info(payments)")

    columns = cursor.fetchall()

    conn.close()

    return "<br>".join(str(col) for col in columns)


@app.route("/check_subscriptions_table")
def check_subscriptions_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("PRAGMA table_info(subscriptions)")

    columns = cursor.fetchall()

    conn.close()

    return "<br>".join(str(col) for col in columns)



# ✅ Update subscriptions table
def update_subscriptions_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Create the new table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS subscriptions_new (

        id INTEGER PRIMARY KEY AUTOINCREMENT,

        username TEXT NOT NULL,

        plan_id INTEGER NOT NULL,

        plan_name TEXT NOT NULL,

        start_date TEXT,

        end_date TEXT,

        payment_reference TEXT,

        payment_status TEXT DEFAULT 'PENDING',

        is_active INTEGER DEFAULT 0,

        amount_paid REAL DEFAULT 0,

        currency TEXT DEFAULT 'NGN',

        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )
    """)

    # Copy existing data
    cursor.execute("""
    INSERT INTO subscriptions_new
    (
        id,
        username,
        plan_id,
        plan_name,
        start_date,
        end_date,
        payment_reference,
        payment_status,
        is_active,
        amount_paid,
        currency,
        created_at
    )
    SELECT
        id,
        email,
        0,
        plan,
        start_date,
        expiry_date,
        payment_reference,
        payment_status,
        is_active,
        amount_paid,
        currency,
        created_at
    FROM subscriptions
    """)

    cursor.execute("DROP TABLE subscriptions")

    cursor.execute(
        "ALTER TABLE subscriptions_new RENAME TO subscriptions"
    )

    conn.commit()
    conn.close()
    


# ✅ Create payments table (it is a temporary route)
@app.route("/create_payments_table")
def create_payments_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS payments (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            username TEXT NOT NULL,

            plan_id INTEGER NOT NULL,

            plan_name TEXT NOT NULL,

            amount REAL NOT NULL,

            duration_days INTEGER NOT NULL,

            transaction_reference TEXT UNIQUE,

            payment_status TEXT DEFAULT 'PENDING',

            payment_method TEXT,

            paid_at TIMESTAMP,

            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

            FOREIGN KEY(plan_id)
                REFERENCES subscription_plans(id)

        )
    """)

    conn.commit()
    conn.close()

    return "payments table created successfully."



# ✅ Bookmark Question
@app.route("/bookmark_question", methods=["POST"])
def bookmark_question():

    if "user" not in session:
        return redirect("/login")

    question_id = request.form["question_id"]
    exam_type = request.form["exam_type"]
    subject = request.form["subject"]

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT OR IGNORE INTO bookmarks
        (
            username,
            question_id,
            exam_type,
            subject
        )
        VALUES (?, ?, ?, ?)
        """,
        (
            session["user"],
            question_id,
            exam_type,
            subject
        )
    )

    conn.commit()
    conn.close()

    return redirect(
        f"/practice_questions/{exam_type}/{subject}"
    )


# ✅ My Bookmarks
@app.route("/my_bookmarks")
def my_bookmarks():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            b.id,
            q.exam_type,
            q.subject,
            q.question_text
        FROM bookmarks b
        JOIN questions q
            ON b.question_id = q.id
        WHERE b.username=?
        ORDER BY b.bookmarked_at DESC
        """,
        (
            session["user"],
        )
    )

    bookmarks = cursor.fetchall()

    conn.close()

    return render_template(
        "my_bookmarks.html",
        bookmarks=bookmarks
    )


# ✅ Delete Bookmark
@app.route("/delete_bookmark/<int:bookmark_id>")
def delete_bookmark(bookmark_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        DELETE FROM bookmarks
        WHERE id=?
        AND username=?
        """,
        (
            bookmark_id,
            session["user"]
        )
    )

    conn.commit()
    conn.close()

    return redirect("/my_bookmarks")


# ✅ My Account
@app.route("/profile")
def profile():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get user information
    cursor.execute(
        """
        SELECT
            name,
            email,
            phone,
            status,
            date_joined,
            profile_picture
        FROM users
        WHERE email=?
        """,
        (
            session["user"],
        )
    )

    user = cursor.fetchone()

    # Get subscription information
    cursor.execute(
        """
        SELECT
            plan_name,
            end_date,
            is_active
        FROM subscriptions
        WHERE username=?
        """,
        (
            session["user"],
        )
    )

    subscription = cursor.fetchone()

    conn.close()

    return render_template(
        "profile.html",
        user=user,
        subscription=subscription
    )


# ✅ Edit Profile
@app.route("/edit_profile", methods=["GET", "POST"])
def edit_profile():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        name = request.form["name"].strip()
        phone = request.form["phone"].strip()

        cursor.execute(
            """
            UPDATE users
            SET
                name=?,
                phone=?
            WHERE email=?
            """,
            (
                name,
                phone,
                session["user"]
            )
        )

        conn.commit()

        # Update the session so the dashboard greeting changes immediately
        session["name"] = name

        conn.close()

        return redirect("/profile")

    cursor.execute(
        """
        SELECT
            name,
            email,
            phone
        FROM users
        WHERE email=?
        """,
        (
            session["user"],
        )
    )

    user = cursor.fetchone()

    conn.close()

    return render_template(
        "edit_profile.html",
        user=user
    )


# ✅ Change Password
@app.route("/change_password", methods=["GET", "POST"])
def change_password():

    if "user" not in session:
        return redirect("/login")

    message = None
    alert_type = "danger"

    if request.method == "POST":

        current_password = request.form["current_password"]
        new_password = request.form["new_password"]
        confirm_password = request.form["confirm_password"]

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT password
            FROM users
            WHERE email=?
            """,
            (
                session["user"],
            )
        )

        row = cursor.fetchone()

        if not row:

            conn.close()
            return redirect("/logout")

        stored_password = row[0]

        # Verify current password
        if not check_password_hash(
            stored_password,
            current_password
        ):

            message = "Current password is incorrect."

        elif new_password != confirm_password:

            message = "New passwords do not match."

        elif len(new_password) < 6:

            message = "Password must be at least 6 characters."

        else:

            hashed_password = generate_password_hash(
                new_password
            )

            cursor.execute(
                """
                UPDATE users
                SET password=?
                WHERE email=?
                """,
                (
                    hashed_password,
                    session["user"]
                )
            )

            conn.commit()

            conn.close()

            # Clear the current session
            session.clear()

            # Redirect to login with a success message
            return redirect("/login?password_changed=1")

        conn.close()

    return render_template(
        "change_password.html",
        message=message,
        alert_type=alert_type
    )


    
# ✅ Delete Subject Combination
@app.route(
    "/delete_subject_combination/<int:combination_id>",
    methods=["GET", "POST"]
)
def delete_subject_combination(combination_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if request.method == "POST":

        cursor.execute(
            """
            DELETE FROM post_utme_course_subjects
            WHERE id=?
            """,
            (combination_id,)
        )

        conn.commit()
        conn.close()

        return redirect(
            "/manage_subject_combinations"
        )

    cursor.execute(
        """
        SELECT *
        FROM post_utme_course_subjects
        WHERE id=?
        """,
        (combination_id,)
    )

    combination = cursor.fetchone()

    conn.close()

    return render_template(
        "delete_subject_combination.html",
        combination=combination
    )


# ✅ Manage Students
@app.route("/manage_students")
def manage_students():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get(
        "search",
        ""
    ).strip()

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    if search:

        cursor.execute(
            """
            SELECT *
            FROM users
            WHERE
                name LIKE ?
                OR email LIKE ?
                OR phone LIKE ?
            ORDER BY id DESC
            """,
            (
                f"%{search}%",
                f"%{search}%",
                f"%{search}%"
            )
        )

    else:

        cursor.execute(
            """
            SELECT *
            FROM users
            ORDER BY id DESC
            """
        )

    students = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_students.html",
        students=students,
        search=search
    )


# ✅ View Student
@app.route("/view_student/<int:user_id>")
def view_student(user_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Student information
    cursor.execute("""
        SELECT *
        FROM users
        WHERE id=?
    """, (user_id,))

    student = cursor.fetchone()

    if not student:
        conn.close()
        return "Student not found."

    # Student results
    cursor.execute("""
        SELECT *
        FROM results
        WHERE username=?
        ORDER BY id DESC
    """, (student["email"],))

    results = cursor.fetchall()
    
    total_exams = len(results)

    if total_exams > 0:

        percentages = [r["percentage"] for r in results]

        average_score = round(sum(percentages) / total_exams)

        highest_score = max(percentages)

        lowest_score = min(percentages)

        pass_count = sum(
            1 for r in results
            if r["status"] == "PASS"
        )

        average_count = sum(
            1 for r in results
            if r["status"] == "AVERAGE"
        )

        fail_count = sum(
            1 for r in results
            if r["status"] == "FAIL"
        )

    else:

        average_score = 0
        highest_score = 0
        lowest_score = 0

        pass_count = 0
        average_count = 0
        fail_count = 0

    conn.close()

    return render_template(
    "view_student.html",
    student=student,
    results=results,

    total_exams=total_exams,
    average_score=average_score,
    highest_score=highest_score,
    lowest_score=lowest_score,

    pass_count=pass_count,
    average_count=average_count,
    fail_count=fail_count
)


# ✅ Suspend Student
@app.route("/suspend_student/<int:user_id>")
def suspend_student(user_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        UPDATE users
        SET
            status='SUSPENDED',
            is_active=0
        WHERE id=?
        """,
        (user_id,)
    )

    conn.commit()
    conn.close()

    return redirect(
        f"/view_student/{user_id}"
    )


# ✅ Activate Student
@app.route("/activate_student/<int:user_id>")
def activate_student(user_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        UPDATE users
        SET
            status='ACTIVE',
            is_active=1
        WHERE id=?
        """,
        (user_id,)
    )

    conn.commit()
    conn.close()

    return redirect(
        f"/view_student/{user_id}"
    )


# ✅ Delete Student Confirmation
@app.route("/delete_student/<int:user_id>")
def delete_student(user_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM users
        WHERE id=?
        """,
        (user_id,)
    )

    student = cursor.fetchone()

    conn.close()

    if not student:
        return "Student not found."

    return render_template(
        "delete_student.html",
        student=student
    )


# ✅ Confirm Delete Student
@app.route("/confirm_delete_student/<int:user_id>")
def confirm_delete_student(user_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        DELETE FROM users
        WHERE id=?
        """,
        (user_id,)
    )

    conn.commit()
    conn.close()

    return redirect("/manage_students")


# ✅ Manage Results
@app.route("/manage_results")
def manage_results():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get("search", "").strip()

    exam_type = request.args.get("exam_type", "").strip()

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    conditions = []
    params = []

    if search:

        conditions.append(
            "(username LIKE ? OR exam_name LIKE ?)"
        )

        params.extend([
            f"%{search}%",
            f"%{search}%"
        ])

    if exam_type:

        conditions.append(
            "exam_type=?"
        )

        params.append(exam_type)

    where_clause = ""

    if conditions:
        where_clause = (
            "WHERE " + " AND ".join(conditions)
        )

    query = f"""
        SELECT *
        FROM results
        {where_clause}
        ORDER BY id DESC
    """

    cursor.execute(query, params)

    results = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_results.html",
        results=results,
        search=search,
        exam_type=exam_type
    )


# ✅ View Result
@app.route("/view_result/<int:result_id>")
def view_result(result_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM results
        WHERE id=?
        """,
        (result_id,)
    )

    result = cursor.fetchone()

    conn.close()

    if not result:
        return "Result not found."

    return render_template(
        "view_result.html",
        result=result
    )


# ✅ Review Result Answers
@app.route("/review_result/<int:result_id>")
def review_result(result_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM review_answers
        WHERE result_id=?
        """,
        (result_id,)
    )

    answers = cursor.fetchall()

    conn.close()

    return render_template(
        "review_result.html",
        answers=answers,
        result_id=result_id
    )
    
    
# ==========================================================
# START / RESUME FULL JAMB EXAM
#
# DATABASE IS THE SOURCE OF TRUTH
# ==========================================================

@app.route("/start_jamb/<course>")
def start_jamb(course):

    # ======================================================
    # CHECK LOGIN
    # ======================================================

    if "user" not in session:
        return redirect("/login")

    username = session["user"]

    # ======================================================
    # GET SUBJECTS FOR SELECTED COURSE
    # ======================================================

    subjects = JAMB_COURSES.get(course)

    if not subjects:
        return "Invalid course selected"

    # ======================================================
    # CHECK FOR EXISTING UNFINISHED JAMB EXAM
    # ======================================================

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            subject,
            q_index,
            score,
            total_answered,
            start_time,
            exam_duration,
            current_subject_index
        FROM exam_progress
        WHERE username=?
        AND exam_type='JAMB'
        AND status='IN_PROGRESS'
        ORDER BY id DESC
        LIMIT 1
        """,
        (username,)
    )

    unfinished = cursor.fetchone()

    conn.close()

    # ======================================================
    # UNFINISHED EXAM EXISTS
    #
    # DO NOT RESUME AUTOMATICALLY.
    #
    # SHOW THE STUDENT A CHOICE.
    # ======================================================

    if unfinished:

        (
            progress_id,
            saved_subject,
            saved_q_index,
            saved_score,
            saved_total_answered,
            saved_start_time,
            saved_exam_duration,
            saved_subject_index
        ) = unfinished

        # --------------------------------------------------
        # Find the course belonging to the unfinished exam
        # --------------------------------------------------

        unfinished_course = None

        for course_name, course_subjects in JAMB_COURSES.items():

            if saved_subject in course_subjects:

                unfinished_course = course_name

                break

        # --------------------------------------------------
        # Safety fallback
        # --------------------------------------------------

        if not unfinished_course:

            unfinished_course = "Unknown"

        # --------------------------------------------------
        # Calculate question number
        # --------------------------------------------------

        question_number = (
            (saved_q_index or 0) + 1
        )

        # --------------------------------------------------
        # Store temporary information so the
        # confirmation buttons know which exam
        # the student is dealing with.
        # --------------------------------------------------

        session["pending_jamb_course"] = course

        session["pending_progress_id"] = progress_id

        # --------------------------------------------------
        # DEBUG
        # --------------------------------------------------

        print("\n====================================")
        print("UNFINISHED JAMB EXAM FOUND")
        print("SELECTED COURSE:", course)
        print("UNFINISHED COURSE:", unfinished_course)
        print("SUBJECT:", saved_subject)
        print("QUESTION:", question_number)
        print("SCORE:", saved_score)
        print("TOTAL ANSWERED:", saved_total_answered)
        print("PROGRESS ID:", progress_id)
        print("====================================\n")

        # --------------------------------------------------
        # SHOW RESUME / NEW EXAM CHOICE
        # --------------------------------------------------

        return render_template(
            "resume_exam.html",

            selected_course=course,

            unfinished_course=unfinished_course,

            saved_subject=saved_subject,

            question_number=question_number,

            score=saved_score or 0,

            total_answered=saved_total_answered or 0
        )

    # ======================================================
    # NO UNFINISHED EXAM
    #
    # START BRAND-NEW EXAM
    # ======================================================

    if not can_take_exam(username):
        return redirect("/subscribe")

    # ------------------------------------------------------
    # Initialize session
    # ------------------------------------------------------

    session["exam_type"] = "JAMB"

    session["course"] = course

    session["subjects"] = subjects

    session["subject_index"] = 0

    session["subject"] = subjects[0]

    session["q_index"] = 0

    session["score"] = 0

    session["total_answered"] = 0

    session["subject_scores"] = {}

    session["subject_totals"] = {}

    session["review_answers"] = []

    session["start_time"] = time.time()

    session["exam_duration"] = 120 * 60

    # ------------------------------------------------------
    # Clear temporary data
    # ------------------------------------------------------

    session.pop(
        "current_question_id",
        None
    )

    session.pop(
        "question_source",
        None
    )

    for key in list(session.keys()):

        if key.startswith("shuffled_"):

            session.pop(
                key,
                None
            )

    # ======================================================
    # FIRST SUBJECT
    # ======================================================

    first_subject = subjects[0]

    # ======================================================
    # CREATE PROGRESS RECORD
    # ======================================================

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO exam_progress
        (
            username,
            exam_type,
            subject,
            q_index,
            score,
            total_answered,
            start_time,
            exam_duration,
            status,
            current_subject_index
        )
        VALUES
        (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            username,
            "JAMB",
            first_subject,
            0,
            0,
            0,
            session["start_time"],
            session["exam_duration"],
            "IN_PROGRESS",
            0
        )
    )

    conn.commit()
    conn.close()

    # ======================================================
    # DEBUG
    # ======================================================

    print("\n====================================")
    print("STARTING NEW JAMB EXAM")
    print("COURSE:", course)
    print("FIRST SUBJECT:", first_subject)
    print("SUBJECT INDEX:", 0)
    print("QUESTION INDEX:", 0)
    print("====================================\n")

    # ======================================================
    # START FIRST SUBJECT
    # ======================================================

    return redirect(
        f"/question/JAMB/{first_subject}"
    )
    
    
    
    
    
# ✅ Subject selection
@app.route("/subjects/<exam_type>")
def subjects(exam_type):

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        "SELECT DISTINCT subject FROM questions WHERE exam_type=?",
        (exam_type,)
    )

    subjects = cursor.fetchall()

    conn.close()

    return render_template(
        "subjects.html",
        subjects=subjects,
        exam_type=exam_type
    )


# ✅ Admin login
@app.route("/admin_login", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        if email == ADMIN_EMAIL and password == ADMIN_PASSWORD:
            
            session.clear()

            session["admin"] = email

            return redirect("/admin")

        else:
            return "Invalid admin credentials"

    return render_template("admin_login.html")

    
# ✅ Admin Dashboard
@app.route("/admin")
def admin():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Students
    cursor.execute("SELECT COUNT(*) FROM users")
    total_students = cursor.fetchone()[0]

    cursor.execute(
        "SELECT COUNT(*) FROM users WHERE is_active=1"
    )
    active_students = cursor.fetchone()[0]

    cursor.execute(
        "SELECT COUNT(*) FROM users WHERE is_active=0"
    )
    suspended_students = cursor.fetchone()[0]

    # WAEC Questions
    cursor.execute(
        "SELECT COUNT(*) FROM questions WHERE exam_type='WAEC'"
    )
    waec_questions = cursor.fetchone()[0]

    # JAMB Questions
    cursor.execute(
        "SELECT COUNT(*) FROM questions WHERE exam_type='JAMB'"
    )
    jamb_questions = cursor.fetchone()[0]

    # POST-UTME Questions
    cursor.execute(
        "SELECT COUNT(*) FROM post_utme_questions"
    )
    post_utme_questions = cursor.fetchone()[0]

    # Results
    cursor.execute(
        "SELECT COUNT(*) FROM results"
    )
    total_results = cursor.fetchone()[0]
    
    # Overall Performance Statistics

    cursor.execute("""
        SELECT
            COUNT(*),
            AVG(percentage),
            MAX(percentage),
            MIN(percentage)
        FROM results
    """)

    stats = cursor.fetchone()

    total_exams_taken = stats[0] if stats[0] else 0

    average_score = round(stats[1]) if stats[1] else 0

    highest_score = stats[2] if stats[2] else 0

    lowest_score = stats[3] if stats[3] else 0


    # Pass Rate

    cursor.execute("""
        SELECT COUNT(*)
        FROM results
        WHERE status='PASS'
    """)

    passed = cursor.fetchone()[0]

    if total_exams_taken > 0:
        pass_rate = round((passed / total_exams_taken) * 100)
    else:
        pass_rate = 0

    conn.close()

    return render_template(
    "admin_dashboard.html",

    total_students=total_students,
    active_students=active_students,
    suspended_students=suspended_students,

    waec_questions=waec_questions,
    jamb_questions=jamb_questions,
    post_utme_questions=post_utme_questions,

    total_results=total_results,

    total_exams_taken=total_exams_taken,
    average_score=average_score,
    highest_score=highest_score,
    lowest_score=lowest_score,
    pass_rate=pass_rate
)


# ✅ View All Results (Admin)
@app.route("/admin/results")
def admin_results():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            username,
            exam_type,
            exam_name,
            score,
            total,
            percentage,
            duration,
            status,
            verification_code,
            date_taken
        FROM results
        ORDER BY id DESC
    """)

    results = cursor.fetchall()

    conn.close()

    return render_template(
        "admin_results.html",
        results=results
    )

    

# ✅ WAEC/JAMB Question Management
@app.route("/admin_questions", methods=["GET", "POST"])
def admin_questions():

    if "admin" not in session:
        return redirect("/admin_login")

    if request.method == "POST":

        exam_type = request.form["exam_type"]
        subject = request.form["subject"]
        question = request.form["question"]

        option_a = request.form["option_a"]
        option_b = request.form["option_b"]
        option_c = request.form["option_c"]
        option_d = request.form["option_d"]

        correct_answer = request.form["correct_answer"]
        explanation = request.form["explanation"]

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT *
            FROM questions
            WHERE exam_type=?
            AND subject=?
            AND question_text=?
            """,
            (
                exam_type,
                subject,
                question
            )
        )

        existing_question = cursor.fetchone()

        if existing_question:

            conn.close()

            return "Question already exists!"

        cursor.execute(
            """
            INSERT INTO questions
            (
                exam_type,
                subject,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                exam_type,
                subject,
                question,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation
            )
        )

        conn.commit()
        conn.close()

        return "Question Added Successfully!"

    return render_template("admin.html")



# ==========================================
# ADD QUESTION (V2)
# ==========================================

@app.route("/admin_questions_v2", methods=["GET", "POST"])
def admin_questions_v2():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # ==========================================
    # SAVE QUESTION
    # ==========================================

    if request.method == "POST":

        exam_type_id = request.form.get("exam_type_id", "").strip()
        subject_id = request.form.get("subject_id", "").strip()
        topic_id = request.form.get("topic_id", "").strip()

        question_type = request.form.get(
            "question_type",
            "Objective"
        ).strip()

        difficulty = request.form.get(
            "difficulty",
            "Medium"
        ).strip()

        question_text = request.form.get(
            "question_text",
            ""
        ).strip()

        option_a = request.form.get(
            "option_a",
            ""
        ).strip()

        option_b = request.form.get(
            "option_b",
            ""
        ).strip()

        option_c = request.form.get(
            "option_c",
            ""
        ).strip()

        option_d = request.form.get(
            "option_d",
            ""
        ).strip()

        correct_answer = request.form.get(
            "correct_answer",
            ""
        ).strip()

        explanation = request.form.get(
            "explanation",
            ""
        ).strip()

        # ==========================================
        # VALIDATION
        # ==========================================

        if not exam_type_id or not subject_id or not topic_id:

            conn.close()

            flash(
                "⚠️ Please select Exam Type, Subject and Topic.",
                "warning"
            )

            return redirect("/admin_questions_v2")

        if not question_text:

            conn.close()

            flash(
                "⚠️ Question text is required.",
                "warning"
            )

            return redirect("/admin_questions_v2")

        if not option_a or not option_b or not option_c or not option_d:

            conn.close()

            flash(
                "⚠️ All four answer options are required.",
                "warning"
            )

            return redirect("/admin_questions_v2")

        if correct_answer not in ["A", "B", "C", "D"]:

            conn.close()

            flash(
                "⚠️ Please select a valid correct answer.",
                "warning"
            )

            return redirect("/admin_questions_v2")

        # ==========================================
        # VERIFY EXAM TYPE + SUBJECT RELATIONSHIP
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM subjects
            WHERE id=?
            AND exam_type_id=?
            AND status='Active'
        """, (
            subject_id,
            exam_type_id
        ))

        valid_subject = cursor.fetchone()

        if not valid_subject:

            conn.close()

            flash(
                "❌ Invalid Subject for the selected Exam Type.",
                "danger"
            )

            return redirect("/admin_questions_v2")

        # ==========================================
        # VERIFY TOPIC + SUBJECT RELATIONSHIP
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM topics
            WHERE id=?
            AND subject_id=?
            AND status='Active'
        """, (
            topic_id,
            subject_id
        ))

        valid_topic = cursor.fetchone()

        if not valid_topic:

            conn.close()

            flash(
                "❌ Invalid Topic for the selected Subject.",
                "danger"
            )

            return redirect("/admin_questions_v2")

        # ==========================================
        # DUPLICATE QUESTION PROTECTION
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM questions_v2
            WHERE exam_type_id=?
            AND subject_id=?
            AND topic_id=?
            AND question_text=?
        """, (
            exam_type_id,
            subject_id,
            topic_id,
            question_text
        ))

        existing_question = cursor.fetchone()

        if existing_question:

            conn.close()

            flash(
                "⚠️ Question already exists.",
                "warning"
            )

            return redirect("/admin_questions_v2")

        # ==========================================
        # INSERT QUESTION
        # ==========================================

        try:

            cursor.execute("""
                INSERT INTO questions_v2
                (
                    exam_type_id,
                    subject_id,
                    topic_id,
                    question_type,
                    difficulty,
                    question_text,
                    option_a,
                    option_b,
                    option_c,
                    option_d,
                    correct_answer,
                    explanation,
                    status
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                exam_type_id,
                subject_id,
                topic_id,
                question_type,
                difficulty,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation,
                "Active"
            ))

            conn.commit()
            conn.close()

            flash(
                "✅ Question added successfully.",
                "success"
            )

            return redirect("/admin_questions_v2")

        except Exception as e:

            conn.rollback()
            conn.close()

            print("QUESTION V2 ERROR:", e)

            flash(
                "❌ Failed to save question.",
                "danger"
            )

            return redirect("/admin_questions_v2")

    # ==========================================
    # LOAD EXAM TYPES
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            exam_name
        FROM exam_types
        WHERE status='Active'
        ORDER BY exam_name
    """)

    exam_types = cursor.fetchall()

    conn.close()

    return render_template(
        "admin_questions_v2.html",
        exam_types=exam_types
    )



# ==========================================
# MANAGE QUESTIONS V2
# ==========================================

@app.route("/manage_questions_v2")
def manage_questions_v2():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get(
        "search",
        ""
    ).strip()

    exam_type_id = request.args.get(
        "exam_type_id",
        ""
    ).strip()

    subject_id = request.args.get(
        "subject_id",
        ""
    ).strip()

    topic_id = request.args.get(
        "topic_id",
        ""
    ).strip()

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # ==========================================
    # VALIDATE EXAM TYPE
    # ==========================================

    if exam_type_id:

        cursor.execute("""
            SELECT id
            FROM exam_types
            WHERE id=?
            AND status='Active'
        """, (exam_type_id,))

        valid_exam = cursor.fetchone()

        if not valid_exam:

            exam_type_id = ""

            flash(
                "⚠️ Invalid exam type filter.",
                "warning"
            )

    # ==========================================
    # VALIDATE SUBJECT
    # ==========================================

    if subject_id:

        if exam_type_id:

            cursor.execute("""
                SELECT id
                FROM subjects
                WHERE id=?
                AND exam_type_id=?
                AND status='Active'
            """, (
                subject_id,
                exam_type_id
            ))

        else:

            cursor.execute("""
                SELECT id
                FROM subjects
                WHERE id=?
                AND status='Active'
            """, (
                subject_id,
            ))

        valid_subject = cursor.fetchone()

        if not valid_subject:

            subject_id = ""

            topic_id = ""

            flash(
                "⚠️ Invalid subject filter for the selected exam.",
                "warning"
            )

    # ==========================================
    # VALIDATE TOPIC
    # ==========================================

    if topic_id:

        if subject_id:

            cursor.execute("""
                SELECT id
                FROM topics
                WHERE id=?
                AND subject_id=?
                AND status='Active'
            """, (
                topic_id,
                subject_id
            ))

        else:

            cursor.execute("""
                SELECT id
                FROM topics
                WHERE id=?
                AND status='Active'
            """, (
                topic_id,
            ))

        valid_topic = cursor.fetchone()

        if not valid_topic:

            topic_id = ""

            flash(
                "⚠️ Invalid topic filter for the selected subject.",
                "warning"
            )

    # ==========================================
    # LOAD QUESTIONS
    # ==========================================

    query = """
        SELECT

            q.id,

            q.question_text,

            q.question_type,

            q.difficulty,

            q.correct_answer,

            q.status,

            q.created_at,

            e.exam_name,

            s.subject_name,

            t.topic_name

        FROM questions_v2 q

        JOIN exam_types e
            ON q.exam_type_id = e.id

        JOIN subjects s
            ON q.subject_id = s.id

        JOIN topics t
            ON q.topic_id = t.id

        WHERE 1=1
    """

    params = []

    # ==========================================
    # SEARCH
    # ==========================================

    if search:

        query += """
            AND q.question_text LIKE ?
        """

        params.append(
            f"%{search}%"
        )

    # ==========================================
    # EXAM TYPE FILTER
    # ==========================================

    if exam_type_id:

        query += """
            AND q.exam_type_id=?
        """

        params.append(
            exam_type_id
        )

    # ==========================================
    # SUBJECT FILTER
    # ==========================================

    if subject_id:

        query += """
            AND q.subject_id=?
        """

        params.append(
            subject_id
        )

    # ==========================================
    # TOPIC FILTER
    # ==========================================

    if topic_id:

        query += """
            AND q.topic_id=?
        """

        params.append(
            topic_id
        )

    # ==========================================
    # ORDER
    # ==========================================

    query += """
        ORDER BY q.id DESC
    """

    cursor.execute(
        query,
        params
    )

    questions = cursor.fetchall()

    # ==========================================
    # LOAD ACTIVE EXAM TYPES
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            exam_name
        FROM exam_types
        WHERE status='Active'
        ORDER BY exam_name
    """)

    exam_types = cursor.fetchall()

    # ==========================================
    # LOAD ACTIVE SUBJECTS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            subject_name,
            exam_type_id
        FROM subjects
        WHERE status='Active'
        ORDER BY subject_name
    """)

    subjects = cursor.fetchall()

    # ==========================================
    # LOAD ACTIVE TOPICS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            topic_name,
            subject_id
        FROM topics
        WHERE status='Active'
        ORDER BY topic_name
    """)

    topics = cursor.fetchall()

    conn.close()

    return render_template(
        "manage_questions_v2.html",

        questions=questions,

        exam_types=exam_types,

        subjects=subjects,

        topics=topics,

        search=search,

        selected_exam_type=exam_type_id,

        selected_subject=subject_id,

        selected_topic=topic_id
    )
    
    

# ==========================================
# QUESTION BANK V2
# WITH FILTERING + PAGINATION
# + SUMMARY STATISTICS
# ==========================================

@app.route("/question_bank_v2")
def question_bank_v2():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # ==========================================
    # GET FILTER VALUES
    # ==========================================

    exam_type_id = request.args.get(
        "exam_type_id",
        ""
    ).strip()

    subject_id = request.args.get(
        "subject_id",
        ""
    ).strip()

    topic_id = request.args.get(
        "topic_id",
        ""
    ).strip()

    question_type = request.args.get(
        "question_type",
        ""
    ).strip()

    difficulty = request.args.get(
        "difficulty",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    search = request.args.get(
        "search",
        ""
    ).strip()


    # ==========================================
    # PAGINATION
    # ==========================================

    try:

        page = int(
            request.args.get(
                "page",
                1
            )
        )

    except ValueError:

        page = 1


    if page < 1:

        page = 1


    # Number of questions per page

    per_page = 25


    # ==========================================
    # BUILD BASE QUERY
    # ==========================================

    base_query = """
        FROM questions_v2 q

        LEFT JOIN exam_types e
            ON q.exam_type_id = e.id

        LEFT JOIN subjects s
            ON q.subject_id = s.id

        LEFT JOIN topics t
            ON q.topic_id = t.id

        WHERE 1=1
    """

    params = []


    # ==========================================
    # EXAM TYPE FILTER
    # ==========================================

    if exam_type_id:

        base_query += """
            AND q.exam_type_id=?
        """

        params.append(
            exam_type_id
        )


    # ==========================================
    # SUBJECT FILTER
    # ==========================================

    if subject_id:

        base_query += """
            AND q.subject_id=?
        """

        params.append(
            subject_id
        )


    # ==========================================
    # TOPIC FILTER
    # ==========================================

    if topic_id:

        base_query += """
            AND q.topic_id=?
        """

        params.append(
            topic_id
        )


    # ==========================================
    # QUESTION TYPE FILTER
    # ==========================================

    if question_type:

        base_query += """
            AND q.question_type=?
        """

        params.append(
            question_type
        )


    # ==========================================
    # DIFFICULTY FILTER
    # ==========================================

    if difficulty:

        base_query += """
            AND q.difficulty=?
        """

        params.append(
            difficulty
        )


    # ==========================================
    # STATUS FILTER
    # ==========================================

    if status:

        base_query += """
            AND q.status=?
        """

        params.append(
            status
        )


    # ==========================================
    # SEARCH FILTER
    # ==========================================

    if search:

        base_query += """
            AND (
                q.question_text LIKE ?
                OR s.subject_name LIKE ?
                OR t.topic_name LIKE ?
            )
        """

        search_value = f"%{search}%"

        params.extend([
            search_value,
            search_value,
            search_value
        ])


    # ==========================================
    # TOTAL QUESTIONS
    # ==========================================

    cursor.execute(
        "SELECT COUNT(*) " + base_query,
        params
    )

    total_questions = cursor.fetchone()[0]


    # ==========================================
    # ACTIVE QUESTIONS
    #
    # Uses the SAME filters
    # ==========================================

    active_query = (
        "SELECT COUNT(*) "
        + base_query
        + " AND q.status='Active'"
    )

    cursor.execute(
        active_query,
        params
    )

    active_questions = cursor.fetchone()[0]


    # ==========================================
    # INACTIVE QUESTIONS
    #
    # Uses the SAME filters
    # ==========================================

    inactive_query = (
        "SELECT COUNT(*) "
        + base_query
        + " AND q.status='Inactive'"
    )

    cursor.execute(
        inactive_query,
        params
    )

    inactive_questions = cursor.fetchone()[0]


    # ==========================================
    # CALCULATE PAGINATION
    # ==========================================

    total_pages = (
        (total_questions + per_page - 1)
        // per_page
    )


    # Make sure page is valid

    if total_pages > 0 and page > total_pages:

        page = total_pages


    if total_pages == 0:

        page = 1


    offset = (
        (page - 1)
        * per_page
    )


    # ==========================================
    # GET QUESTIONS FOR CURRENT PAGE
    # ==========================================

    question_query = """
        SELECT

            q.id,

            q.exam_type_id,

            q.subject_id,

            q.topic_id,

            q.question_type,

            q.difficulty,

            q.question_text,

            q.status,

            e.exam_name,

            s.subject_name,

            t.topic_name

    """ + base_query + """

        ORDER BY q.id DESC

        LIMIT ? OFFSET ?
    """


    question_params = list(params)

    question_params.extend([
        per_page,
        offset
    ])


    cursor.execute(
        question_query,
        question_params
    )

    questions = cursor.fetchall()


    # ==========================================
    # LOAD EXAM TYPES
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            exam_name

        FROM exam_types

        WHERE status='Active'

        ORDER BY exam_name
    """)

    exam_types = cursor.fetchall()


    # ==========================================
    # LOAD ACTIVE SUBJECTS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            subject_name,
            exam_type_id

        FROM subjects

        WHERE status='Active'

        ORDER BY subject_name
    """)

    subjects = cursor.fetchall()


    # ==========================================
    # LOAD ACTIVE TOPICS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            topic_name,
            subject_id

        FROM topics

        WHERE status='Active'

        ORDER BY topic_name
    """)

    topics = cursor.fetchall()


    # ==========================================
    # CLOSE DATABASE
    # ==========================================

    conn.close()


    # ==========================================
    # RENDER PAGE
    # ==========================================

    return render_template(

        "question_bank_v2.html",

        questions=questions,

        exam_types=exam_types,

        subjects=subjects,

        topics=topics,

        exam_type_id=exam_type_id,

        subject_id=subject_id,

        topic_id=topic_id,

        question_type=question_type,

        difficulty=difficulty,

        status=status,

        search=search,

        page=page,

        per_page=per_page,

        total_questions=total_questions,

        active_questions=active_questions,

        inactive_questions=inactive_questions,

        total_pages=total_pages

    )



# ==========================================
# VIEW QUESTION V2
# ==========================================

@app.route("/view_question_v2/<int:question_id>")
def view_question_v2(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("""
        SELECT

            q.id,

            q.question_text,

            q.option_a,
            q.option_b,
            q.option_c,
            q.option_d,

            q.correct_answer,

            q.explanation,

            q.question_type,

            q.difficulty,

            q.status,

            q.created_at,

            e.exam_name,

            s.subject_name,

            t.topic_name

        FROM questions_v2 q

        INNER JOIN exam_types e
            ON q.exam_type_id = e.id

        INNER JOIN subjects s
            ON q.subject_id = s.id

        INNER JOIN topics t
            ON q.topic_id = t.id

        WHERE q.id = ?

    """, (question_id,))

    question = cursor.fetchone()

    conn.close()

    # ==========================================
    # QUESTION NOT FOUND
    # ==========================================

    if not question:

        flash(
            "❌ Question not found.",
            "danger"
        )

        return redirect("/question_bank_v2")

    return render_template(
        "view_question_v2.html",
        question=question
    )



# ==========================================
# EDIT QUESTION V2
# ==========================================

@app.route(
    "/edit_question_v2/<int:question_id>",
    methods=["GET", "POST"]
)
def edit_question_v2(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # ==========================================
    # GET EXISTING QUESTION
    # ==========================================

    cursor.execute("""
        SELECT *
        FROM questions_v2
        WHERE id=?
    """, (question_id,))

    question = cursor.fetchone()

    if not question:

        conn.close()

        flash(
            "❌ Question not found.",
            "danger"
        )

        return redirect("/question_bank_v2")


    # ==========================================
    # UPDATE QUESTION
    # ==========================================

    if request.method == "POST":

        exam_type_id = request.form.get(
            "exam_type_id",
            ""
        ).strip()

        subject_id = request.form.get(
            "subject_id",
            ""
        ).strip()

        topic_id = request.form.get(
            "topic_id",
            ""
        ).strip()

        question_type = request.form.get(
            "question_type",
            "Objective"
        ).strip()

        difficulty = request.form.get(
            "difficulty",
            "Medium"
        ).strip()

        question_text = request.form.get(
            "question_text",
            ""
        ).strip()

        option_a = request.form.get(
            "option_a",
            ""
        ).strip()

        option_b = request.form.get(
            "option_b",
            ""
        ).strip()

        option_c = request.form.get(
            "option_c",
            ""
        ).strip()

        option_d = request.form.get(
            "option_d",
            ""
        ).strip()

        correct_answer = request.form.get(
            "correct_answer",
            ""
        ).strip().upper()

        explanation = request.form.get(
            "explanation",
            ""
        ).strip()

        status = request.form.get(
            "status",
            "Active"
        ).strip()


        # ==========================================
        # BASIC VALIDATION
        # ==========================================

        if not exam_type_id:

            conn.close()

            flash(
                "⚠️ Please select an Exam Type.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if not subject_id:

            conn.close()

            flash(
                "⚠️ Please select a Subject.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if not topic_id:

            conn.close()

            flash(
                "⚠️ Please select a Topic.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if not question_text:

            conn.close()

            flash(
                "⚠️ Question text is required.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if (
            not option_a
            or not option_b
            or not option_c
            or not option_d
        ):

            conn.close()

            flash(
                "⚠️ All four answer options are required.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if correct_answer not in [
            "A",
            "B",
            "C",
            "D"
        ]:

            conn.close()

            flash(
                "⚠️ Please select a valid correct answer.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if question_type not in [
            "Objective",
            "Theory"
        ]:

            conn.close()

            flash(
                "⚠️ Invalid question type.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if difficulty not in [
            "Easy",
            "Medium",
            "Hard"
        ]:

            conn.close()

            flash(
                "⚠️ Invalid difficulty level.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        if status not in [
            "Active",
            "Inactive"
        ]:

            conn.close()

            flash(
                "⚠️ Invalid question status.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        # ==========================================
        # VERIFY EXAM TYPE
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM exam_types
            WHERE id=?
            AND status='Active'
        """, (
            exam_type_id,
        ))

        valid_exam_type = cursor.fetchone()

        if not valid_exam_type:

            conn.close()

            flash(
                "❌ Invalid or inactive Exam Type.",
                "danger"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        # ==========================================
        # VERIFY SUBJECT
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM subjects
            WHERE id=?
            AND exam_type_id=?
            AND status='Active'
        """, (
            subject_id,
            exam_type_id
        ))

        valid_subject = cursor.fetchone()

        if not valid_subject:

            conn.close()

            flash(
                "❌ Invalid Subject for the selected Exam Type.",
                "danger"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        # ==========================================
        # VERIFY TOPIC
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM topics
            WHERE id=?
            AND subject_id=?
            AND status='Active'
        """, (
            topic_id,
            subject_id
        ))

        valid_topic = cursor.fetchone()

        if not valid_topic:

            conn.close()

            flash(
                "❌ Invalid Topic for the selected Subject.",
                "danger"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        # ==========================================
        # DUPLICATE PROTECTION
        # ==========================================

        cursor.execute("""
            SELECT id
            FROM questions_v2
            WHERE exam_type_id=?
            AND subject_id=?
            AND topic_id=?
            AND question_text=?
            AND id != ?
        """, (
            exam_type_id,
            subject_id,
            topic_id,
            question_text,
            question_id
        ))

        duplicate = cursor.fetchone()

        if duplicate:

            conn.close()

            flash(
                "⚠️ Another question with the same "
                "text already exists.",
                "warning"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


        # ==========================================
        # UPDATE DATABASE
        # ==========================================

        try:

            cursor.execute("""
                UPDATE questions_v2

                SET

                    exam_type_id=?,

                    subject_id=?,

                    topic_id=?,

                    question_type=?,

                    difficulty=?,

                    question_text=?,

                    option_a=?,

                    option_b=?,

                    option_c=?,

                    option_d=?,

                    correct_answer=?,

                    explanation=?,

                    status=?

                WHERE id=?

            """, (
                exam_type_id,
                subject_id,
                topic_id,
                question_type,
                difficulty,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation,
                status,
                question_id
            ))


            conn.commit()
            conn.close()


            flash(
                "✅ Question updated successfully.",
                "success"
            )


            return redirect(
                f"/view_question_v2/{question_id}"
            )


        except Exception as e:

            conn.rollback()
            conn.close()

            print(
                "EDIT QUESTION V2 ERROR:",
                e
            )

            flash(
                "❌ Failed to update question.",
                "danger"
            )

            return redirect(
                f"/edit_question_v2/{question_id}"
            )


    # ==========================================
    # LOAD EXAM TYPES
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            exam_name
        FROM exam_types
        WHERE status='Active'
        ORDER BY exam_name
    """)

    exam_types = cursor.fetchall()


    # ==========================================
    # LOAD ACTIVE SUBJECTS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            subject_name,
            exam_type_id
        FROM subjects
        WHERE status='Active'
        ORDER BY subject_name
    """)

    subjects = cursor.fetchall()


    # ==========================================
    # LOAD ACTIVE TOPICS
    # ==========================================

    cursor.execute("""
        SELECT
            id,
            topic_name,
            subject_id
        FROM topics
        WHERE status='Active'
        ORDER BY topic_name
    """)

    topics = cursor.fetchall()


    conn.close()


    return render_template(
        "edit_question_v2.html",

        question=question,

        exam_types=exam_types,

        subjects=subjects,

        topics=topics
    )


# ==========================================
# ACTIVATE / DEACTIVATE QUESTION V2
# ==========================================

@app.route("/toggle_question_v2/<int:question_id>")
def toggle_question_v2(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # ==========================================
    # GET CURRENT STATUS
    # ==========================================

    cursor.execute("""
        SELECT id, status
        FROM questions_v2
        WHERE id=?
    """, (question_id,))

    question = cursor.fetchone()

    if not question:

        conn.close()

        flash(
            "❌ Question not found.",
            "danger"
        )

        return redirect("/question_bank_v2")


    # ==========================================
    # DETERMINE NEW STATUS
    # ==========================================

    if question["status"] == "Active":

        new_status = "Inactive"

        message = "✅ Question deactivated successfully."

        category = "success"

    else:

        new_status = "Active"

        message = "✅ Question activated successfully."

        category = "success"


    # ==========================================
    # UPDATE STATUS
    # ==========================================

    try:

        cursor.execute("""
            UPDATE questions_v2

            SET status=?

            WHERE id=?

        """, (
            new_status,
            question_id
        ))

        conn.commit()
        conn.close()


        # ======================================
        # FLASH MESSAGE
        # ======================================

        flash(
            message,
            category
        )


        # ======================================
        # RETURN TO QUESTION VIEW
        # ======================================

        return redirect(
            f"/view_question_v2/{question_id}"
        )


    except Exception as e:

        conn.rollback()
        conn.close()

        print(
            "TOGGLE QUESTION V2 ERROR:",
            e
        )

        flash(
            "❌ Failed to change question status.",
            "danger"
        )

        return redirect(
            f"/view_question_v2/{question_id}"
        )


# ==========================================
# BULK QUESTION ACTIONS V2
# ==========================================

@app.route("/bulk_question_action_v2", methods=["POST"])
def bulk_question_action_v2():

    if "admin" not in session:
        return redirect("/admin_login")

    # ==========================================
    # GET SELECTED QUESTION IDS
    # ==========================================

    question_ids = request.form.getlist(
        "question_ids"
    )

    action = request.form.get(
        "action",
        ""
    ).strip()


    # ==========================================
    # CHECK SELECTION
    # ==========================================

    if not question_ids:

        flash(
            "⚠️ Please select at least one question.",
            "warning"
        )

        return redirect(
            request.referrer or
            "/question_bank_v2"
        )


    # ==========================================
    # VALIDATE ACTION
    # ==========================================

    if action not in [
        "activate",
        "deactivate",
        "delete"
    ]:

        flash(
            "❌ Invalid bulk action.",
            "danger"
        )

        return redirect(
            request.referrer or
            "/question_bank_v2"
        )


    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()


    try:

        # ==========================================
        # CLEAN QUESTION IDS
        # ==========================================

        valid_ids = []

        for question_id in question_ids:

            try:

                question_id = int(
                    question_id
                )

                valid_ids.append(
                    question_id
                )

            except ValueError:

                continue


        if not valid_ids:

            conn.close()

            flash(
                "⚠️ No valid questions were selected.",
                "warning"
            )

            return redirect(
                request.referrer or
                "/question_bank_v2"
            )


        # ==========================================
        # CREATE SQL PLACEHOLDERS
        # ==========================================

        placeholders = ",".join(
            ["?"] * len(valid_ids)
        )


        # ==========================================
        # ACTIVATE
        # ==========================================

        if action == "activate":

            cursor.execute(
                f"""
                UPDATE questions_v2

                SET status='Active'

                WHERE id IN ({placeholders})
                """,
                valid_ids
            )

            affected_count = cursor.rowcount

            conn.commit()
            conn.close()

            flash(
                f"🟢 {affected_count} question(s) activated successfully.",
                "success"
            )

            return redirect(
                request.referrer or
                "/question_bank_v2"
            )


        # ==========================================
        # DEACTIVATE
        # ==========================================

        elif action == "deactivate":

            cursor.execute(
                f"""
                UPDATE questions_v2

                SET status='Inactive'

                WHERE id IN ({placeholders})
                """,
                valid_ids
            )

            affected_count = cursor.rowcount

            conn.commit()
            conn.close()

            flash(
                f"🔴 {affected_count} question(s) deactivated successfully.",
                "success"
            )

            return redirect(
                request.referrer or
                "/question_bank_v2"
            )


        # ==========================================
        # DELETE
        # ==========================================

        elif action == "delete":

            cursor.execute(
                f"""
                DELETE FROM questions_v2

                WHERE id IN ({placeholders})
                """,
                valid_ids
            )

            affected_count = cursor.rowcount

            conn.commit()
            conn.close()

            flash(
                f"🗑️ {affected_count} question(s) deleted successfully.",
                "success"
            )

            return redirect(
                request.referrer or
                "/question_bank_v2"
            )


    except Exception as e:

        conn.rollback()
        conn.close()

        print(
            "BULK QUESTION ACTION V2 ERROR:",
            e
        )

        flash(
            "❌ Failed to complete the bulk action.",
            "danger"
        )

        return redirect(
            request.referrer or
            "/question_bank_v2"
        )



# ==========================================
# DELETE QUESTION V2
# ==========================================

@app.route("/delete_question_v2/<int:question_id>", methods=["POST"])
def delete_question_v2(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # ==========================================
    # CHECK QUESTION EXISTS
    # ==========================================

    cursor.execute("""
        SELECT id
        FROM questions_v2
        WHERE id=?
    """, (question_id,))

    question = cursor.fetchone()

    if not question:

        conn.close()

        flash(
            "❌ Question not found.",
            "danger"
        )

        return redirect("/question_bank_v2")


    # ==========================================
    # DELETE QUESTION
    # ==========================================

    try:

        cursor.execute("""
            DELETE FROM questions_v2
            WHERE id=?
        """, (question_id,))

        conn.commit()
        conn.close()


        # ======================================
        # SUCCESS MESSAGE
        # ======================================

        flash(
            "✅ Question deleted successfully.",
            "success"
        )


        return redirect("/question_bank_v2")


    except Exception as e:

        conn.rollback()
        conn.close()

        print(
            "DELETE QUESTION V2 ERROR:",
            e
        )

        flash(
            "❌ Failed to delete question.",
            "danger"
        )

        return redirect(
            f"/view_question_v2/{question_id}"
        )
        


# ==========================================
# BULK IMPORT QUESTIONS V2
# ==========================================

@app.route("/import_questions_v2", methods=["GET", "POST"])
def import_questions_v2():

    if "admin" not in session:
        return redirect("/admin_login")

    if request.method == "POST":

        # ==========================================
        # CHECK FILE
        # ==========================================

        if "excel_file" not in request.files:

            flash(
                "❌ No Excel file was selected.",
                "danger"
            )

            return redirect("/import_questions_v2")

        file = request.files["excel_file"]

        if file.filename == "":

            flash(
                "❌ No Excel file was selected.",
                "danger"
            )

            return redirect("/import_questions_v2")

        # ==========================================
        # CHECK FILE TYPE
        # ==========================================

        if not file.filename.lower().endswith(".xlsx"):

            flash(
                "❌ Only .xlsx Excel files are supported.",
                "danger"
            )

            return redirect("/import_questions_v2")

        # ==========================================
        # OPEN EXCEL FILE
        # ==========================================

        try:

            workbook = load_workbook(
                file,
                read_only=True,
                data_only=True
            )

            worksheet = workbook.active

        except Exception as e:

            flash(
                "❌ Unable to read the Excel file.",
                "danger"
            )

            return redirect("/import_questions_v2")

        # ==========================================
        # DATABASE
        # ==========================================

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        # ==========================================
        # COUNTERS
        # ==========================================

        imported = 0
        duplicates = 0
        errors = 0

        error_messages = []

        # ==========================================
        # READ HEADER
        # ==========================================

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:

            header = next(rows)

        except StopIteration:

            workbook.close()
            conn.close()

            flash(
                "❌ The Excel file is empty.",
                "danger"
            )

            return redirect("/import_questions_v2")

        # ==========================================
        # NORMALIZE HEADERS
        # ==========================================

        headers = []

        for value in header:

            if value is None:

                headers.append("")

            else:

                headers.append(
                    str(value)
                    .strip()
                    .lower()
                )

        # ==========================================
        # REQUIRED COLUMNS
        # ==========================================

        required_columns = [

            "exam type",
            "subject",
            "topic",
            "question",
            "option a",
            "option b",
            "option c",
            "option d",
            "answer",
            "explanation",
            "difficulty"

        ]

        # ==========================================
        # CHECK REQUIRED COLUMNS
        # ==========================================

        missing_columns = []

        for column in required_columns:

            if column not in headers:

                missing_columns.append(column)

        if missing_columns:

            workbook.close()
            conn.close()

            flash(
                "❌ Missing Excel columns: "
                + ", ".join(missing_columns),
                "danger"
            )

            return redirect("/import_questions_v2")

        # ==========================================
        # COLUMN INDEXES
        # ==========================================

        exam_type_index = headers.index(
            "exam type"
        )

        subject_index = headers.index(
            "subject"
        )

        topic_index = headers.index(
            "topic"
        )

        question_index = headers.index(
            "question"
        )

        option_a_index = headers.index(
            "option a"
        )

        option_b_index = headers.index(
            "option b"
        )

        option_c_index = headers.index(
            "option c"
        )

        option_d_index = headers.index(
            "option d"
        )

        answer_index = headers.index(
            "answer"
        )

        explanation_index = headers.index(
            "explanation"
        )

        difficulty_index = headers.index(
            "difficulty"
        )

        # ==========================================
        # PROCESS EACH ROW
        # ==========================================

        for row_number, row in enumerate(
            rows,
            start=2
        ):

            try:

                # ----------------------------------
                # SAFELY GET CELL VALUES
                # ----------------------------------

                def get_value(index):

                    if index >= len(row):

                        return ""

                    value = row[index]

                    if value is None:

                        return ""

                    return str(value).strip()

                exam_type_name = get_value(
                    exam_type_index
                )

                subject_name = get_value(
                    subject_index
                )

                topic_name = get_value(
                    topic_index
                )

                question_text = get_value(
                    question_index
                )

                option_a = get_value(
                    option_a_index
                )

                option_b = get_value(
                    option_b_index
                )

                option_c = get_value(
                    option_c_index
                )

                option_d = get_value(
                    option_d_index
                )

                correct_answer = get_value(
                    answer_index
                ).upper()

                explanation = get_value(
                    explanation_index
                )

                difficulty = get_value(
                    difficulty_index
                )

                # ----------------------------------
                # SKIP COMPLETELY EMPTY ROW
                # ----------------------------------

                if not any(row):

                    continue

                # ----------------------------------
                # REQUIRED DATA
                # ----------------------------------

                if not exam_type_name:

                    raise ValueError(
                        "Exam Type is empty."
                    )

                if not subject_name:

                    raise ValueError(
                        "Subject is empty."
                    )

                if not topic_name:

                    raise ValueError(
                        "Topic is empty."
                    )

                if not question_text:

                    raise ValueError(
                        "Question is empty."
                    )

                if not correct_answer:

                    raise ValueError(
                        "Answer is empty."
                    )

                # ----------------------------------
                # VALIDATE ANSWER
                # ----------------------------------

                if correct_answer not in [
                    "A",
                    "B",
                    "C",
                    "D"
                ]:

                    raise ValueError(
                        "Answer must be A, B, C or D."
                    )

                # ----------------------------------
                # FIND EXAM TYPE
                # ----------------------------------

                cursor.execute("""
                    SELECT id
                    FROM exam_types
                    WHERE LOWER(exam_name)=LOWER(?)
                    AND status='Active'
                """, (
                    exam_type_name,
                ))

                exam_type = cursor.fetchone()

                if not exam_type:

                    raise ValueError(
                        f"Exam Type '{exam_type_name}' "
                        f"was not found or is inactive."
                    )

                exam_type_id = exam_type[0]

                # ----------------------------------
                # FIND SUBJECT
                # ----------------------------------

                cursor.execute("""
                    SELECT id
                    FROM subjects
                    WHERE LOWER(subject_name)=LOWER(?)
                    AND exam_type_id=?
                    AND status='Active'
                """, (
                    subject_name,
                    exam_type_id
                ))

                subject = cursor.fetchone()

                if not subject:

                    raise ValueError(
                        f"Subject '{subject_name}' "
                        f"does not belong to "
                        f"'{exam_type_name}' "
                        f"or is inactive."
                    )

                subject_id = subject[0]

                # ----------------------------------
                # FIND TOPIC
                # ----------------------------------

                cursor.execute("""
                    SELECT id
                    FROM topics
                    WHERE LOWER(topic_name)=LOWER(?)
                    AND subject_id=?
                    AND status='Active'
                """, (
                    topic_name,
                    subject_id
                ))

                topic = cursor.fetchone()

                if not topic:

                    raise ValueError(
                        f"Topic '{topic_name}' "
                        f"does not belong to "
                        f"'{subject_name}' "
                        f"or is inactive."
                    )

                topic_id = topic[0]

                # ----------------------------------
                # CHECK DUPLICATE
                # ----------------------------------

                cursor.execute("""
                    SELECT id
                    FROM questions_v2
                    WHERE subject_id=?
                    AND topic_id=?
                    AND question_text=?
                """, (
                    subject_id,
                    topic_id,
                    question_text
                ))

                existing_question = cursor.fetchone()

                if existing_question:

                    duplicates += 1

                    error_messages.append(
                        f"Row {row_number}: "
                        f"Duplicate question skipped."
                    )

                    continue

                # ----------------------------------
                # NORMALIZE DIFFICULTY
                # ----------------------------------

                difficulty_map = {

                    "easy": "Easy",

                    "medium": "Medium",

                    "hard": "Hard"

                }

                difficulty_value = difficulty_map.get(
                    difficulty.lower(),
                    "Medium"
                )

                # ----------------------------------
                # INSERT QUESTION
                # ----------------------------------

                cursor.execute("""
                    INSERT INTO questions_v2
                    (
                        exam_type_id,
                        subject_id,
                        topic_id,
                        question_type,
                        difficulty,
                        question_text,
                        option_a,
                        option_b,
                        option_c,
                        option_d,
                        correct_answer,
                        explanation,
                        status
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    exam_type_id,
                    subject_id,
                    topic_id,
                    "Objective",
                    difficulty_value,
                    question_text,
                    option_a,
                    option_b,
                    option_c,
                    option_d,
                    correct_answer,
                    explanation,
                    "Active"
                ))

                imported += 1

            except Exception as e:

                errors += 1

                error_messages.append(
                    f"Row {row_number}: {str(e)}"
                )

        # ==========================================
        # SAVE IMPORTED QUESTIONS
        # ==========================================

        conn.commit()

        # ==========================================
        # PREPARE IMPORT DETAILS
        # ==========================================

        import_details = "\n".join(
            error_messages
        )

        # ==========================================
        # SAVE IMPORT HISTORY
        # ==========================================

        cursor.execute("""
            INSERT INTO question_import_logs
            (
                filename,
                imported_count,
                skipped_count,
                error_count,
                uploaded_by,
                details
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            file.filename,
            imported,
            duplicates,
            errors,
            session.get("user", "Admin"),
            import_details
        ))

        # ==========================================
        # IMPORTANT:
        # COMMIT IMPORT HISTORY
        # ==========================================

        conn.commit()

        # ==========================================
        # CLOSE RESOURCES
        # ==========================================

        workbook.close()
        conn.close()

        # ==========================================
        # STORE IMPORT RESULTS
        # ==========================================

        session["import_results"] = {

            "imported": imported,

            "duplicates": duplicates,

            "errors": errors,

            "error_messages": error_messages[:50]

        }

        return redirect(
            "/import_questions_v2"
        )

    # ==========================================
    # DISPLAY IMPORT RESULTS
    # ==========================================

    import_results = session.pop(
        "import_results",
        None
    )

    return render_template(
        "import_questions_v2.html",
        import_results=import_results
    )




# ==========================================
# DOWNLOAD QUESTION IMPORT TEMPLATE
# ==========================================

@app.route("/download_question_template")
def download_question_template():

    if "admin" not in session:
        return redirect("/admin_login")

    from openpyxl import Workbook

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Questions"

    # ==========================================
    # HEADERS
    # ==========================================

    headers = [
        "Exam Type",
        "Subject",
        "Topic",
        "Question",
        "Option A",
        "Option B",
        "Option C",
        "Option D",
        "Answer",
        "Explanation",
        "Difficulty"
    ]

    worksheet.append(headers)

    # ==========================================
    # SAMPLE QUESTION
    # ==========================================

    worksheet.append([
        "JAMB",
        "Mathematics",
        "Sets",
        "Which of the following is a subset of {1, 2, 3}?",
        "{1, 2}",
        "{4, 5}",
        "{2, 4}",
        "{5, 6}",
        "A",
        "Both 1 and 2 are elements of the original set, so {1, 2} is a subset.",
        "Easy"
    ])

    # ==========================================
    # COLUMN WIDTHS
    # ==========================================

    widths = {
        "A": 15,
        "B": 25,
        "C": 25,
        "D": 60,
        "E": 30,
        "F": 30,
        "G": 30,
        "H": 30,
        "I": 12,
        "J": 60,
        "K": 15
    }

    for column, width in widths.items():

        worksheet.column_dimensions[column].width = width

    # ==========================================
    # FREEZE HEADER
    # ==========================================

    worksheet.freeze_panes = "A2"

    # ==========================================
    # SAVE TEMPORARILY
    # ==========================================

    file_path = "question_import_template.xlsx"

    workbook.save(file_path)

    # ==========================================
    # SEND FILE
    # ==========================================

    return send_file(
        file_path,
        as_attachment=True,
        download_name="question_import_template.xlsx"
    )



# ==========================================
# QUESTION IMPORT HISTORY
# ==========================================

@app.route("/question_import_history")
def question_import_history():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row

    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            filename,
            imported_count,
            skipped_count,
            error_count,
            uploaded_by,
            uploaded_at
        FROM question_import_logs
        ORDER BY id DESC
    """)

    import_logs = cursor.fetchall()

    conn.close()

    return render_template(
        "question_import_history.html",
        import_logs=import_logs
    )



# ==========================================
# VIEW IMPORT DETAILS
# ==========================================

@app.route("/question_import_details/<int:log_id>")
def question_import_details(log_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row

    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            filename,
            imported_count,
            skipped_count,
            error_count,
            uploaded_by,
            uploaded_at,
            details
        FROM question_import_logs
        WHERE id=?
    """, (log_id,))

    import_log = cursor.fetchone()

    conn.close()

    if not import_log:

        flash(
            "❌ Import record not found.",
            "danger"
        )

        return redirect(
            "/question_import_history"
        )

    return render_template(
        "question_import_details.html",
        import_log=import_log
    )



# ==========================================
# GET SUBJECTS BY EXAM TYPE (AJAX)
# ==========================================

@app.route("/get_subjects/<int:exam_type_id>")
def get_subjects(exam_type_id):

    if "admin" not in session:
        return {"subjects": []}

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            subject_name
        FROM subjects
        WHERE
            exam_type_id=?
            AND status='Active'
        ORDER BY subject_name
    """, (exam_type_id,))

    subjects = cursor.fetchall()

    conn.close()

    return {
        "subjects": [
            dict(subject)
            for subject in subjects
        ]
    }



# ==========================================
# GET TOPICS BY SUBJECT (AJAX)
# ==========================================

@app.route("/get_topics/<int:subject_id>")
def get_topics(subject_id):

    if "admin" not in session:
        return {"topics": []}

    conn = sqlite3.connect("database.db")
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            topic_name
        FROM topics
        WHERE
            subject_id=?
            AND status='Active'
        ORDER BY topic_name
    """, (subject_id,))

    topics = cursor.fetchall()

    conn.close()

    return {
        "topics": [
            dict(topic)
            for topic in topics
        ]
    }

    

# ✅ Download Excel Template
@app.route("/download_template")
def download_template():

    if "admin" not in session:
        return redirect("/admin_login")

    columns = [
        "exam_type",
        "subject",
        "question_text",
        "option_a",
        "option_b",
        "option_c",
        "option_d",
        "correct_answer",
        "explanation"
    ]

    df = pd.DataFrame(columns=columns)

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            index=False,
            sheet_name="Questions"
        )

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="questions_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )



# ==========================================================
# QUESTION PAGE
#
# V2 QUESTION BANK FOR JAMB / WAEC
# POST-UTME REMAINS ON CURRENT SYSTEM
# ==========================================================

@app.route("/question/<exam_type>/<subject>")
def question(exam_type, subject):

    # ======================================================
    # CHECK LOGIN
    # ======================================================

    if "user" not in session:
        return redirect("/login")


    # ======================================================
    # CHECK EXAM TIME
    # ======================================================

    elapsed = (
        time.time()
        - session.get(
            "start_time",
            time.time()
        )
    )

    exam_duration = session.get(
        "exam_duration",
        0
    )

    if exam_duration > 0 and elapsed >= exam_duration:

        print("\n====================================")
        print("EXAM TIME EXPIRED")
        print("EXAM TYPE:", exam_type)
        print("SUBJECT:", subject)
        print("====================================\n")

        return redirect("/result")


    # ======================================================
    # IMPORTANT
    #
    # DO NOT RESET q_index HERE.
    #
    # q_index is controlled by:
    #
    # 1. start_jamb() when starting/resuming
    # 2. next_question() when moving forward
    #
    # ======================================================

    if "q_index" not in session:

        session["q_index"] = 0


    # ======================================================
    # SAVE CURRENT EXAM INFORMATION
    # ======================================================

    session["exam_type"] = exam_type
    session["subject"] = subject


    # ======================================================
    # CURRENT QUESTION INDEX
    # ======================================================

    current_index = session.get(
        "q_index",
        0
    )


    # ======================================================
    # CONNECT DATABASE
    # ======================================================

    conn = sqlite3.connect(
        "database.db"
    )

    cursor = conn.cursor()


    # ======================================================
    # POST-UTME
    #
    # KEEP EXISTING SYSTEM
    # ======================================================

    if exam_type == "POST-UTME":

        cursor.execute(
            """
            SELECT *
            FROM post_utme_questions
            WHERE university_name=?
            AND subject=?
            ORDER BY id
            """,
            (
                session["university"],
                subject
            )
        )

        questions = cursor.fetchall()

        question_source = (
            "post_utme_questions"
        )


    # ======================================================
    # JAMB / WAEC
    #
    # USE QUESTIONS_V2
    # ======================================================

    else:

        cursor.execute(
            """
            SELECT

                q.id,

                e.exam_name,

                s.subject_name,

                q.question_text,

                q.option_a,

                q.option_b,

                q.option_c,

                q.option_d,

                q.correct_answer,

                q.explanation

            FROM questions_v2 q

            INNER JOIN exam_types e
                ON q.exam_type_id = e.id

            INNER JOIN subjects s
                ON q.subject_id = s.id

            WHERE
                e.exam_name=?
                AND s.subject_name=?
                AND s.exam_type_id=e.id
                AND e.status='Active'
                AND s.status='Active'
                AND q.status='Active'

            ORDER BY q.id
            """,
            (
                exam_type,
                subject
            )
        )

        questions = cursor.fetchall()

        question_source = (
            "questions_v2"
        )


    # ======================================================
    # CLOSE DATABASE
    # ======================================================

    conn.close()


    # ======================================================
    # CHECK WHETHER QUESTIONS EXIST
    # ======================================================

    if not questions:

        print("\n====================================")
        print("NO QUESTIONS AVAILABLE")
        print("EXAM TYPE:", exam_type)
        print("SUBJECT:", subject)
        print("QUESTION SOURCE:", question_source)
        print("====================================\n")

        return "No questions available."


    # ======================================================
    # CHECK QUESTION INDEX
    # ======================================================

    if current_index >= len(questions):

        print("\n====================================")
        print("QUESTION INDEX EXCEEDED")
        print("EXAM TYPE:", exam_type)
        print("SUBJECT:", subject)
        print("Q_INDEX:", current_index)
        print("TOTAL QUESTIONS:", len(questions))
        print("====================================\n")

        return redirect("/next_question")


    # ======================================================
    # GET CURRENT QUESTION
    # ======================================================

    q = questions[current_index]


    # ======================================================
    # SAVE CURRENT QUESTION INFORMATION
    # ======================================================

    session["current_question_id"] = q[0]

    session["question_source"] = (
        question_source
    )


    # ======================================================
    # DEBUG INFORMATION
    # ======================================================

    print("\n====================================")
    print("DISPLAYING QUESTION")
    print("====================================")

    print(
        "EXAM TYPE:",
        exam_type
    )

    print(
        "SUBJECT:",
        subject
    )

    print(
        "QUESTION SOURCE:",
        question_source
    )

    print(
        "Q_INDEX:",
        current_index
    )

    print(
        "DISPLAY QUESTION:",
        current_index + 1
    )

    print(
        "TOTAL QUESTIONS:",
        len(questions)
    )

    print(
        "QUESTION ID:",
        q[0]
    )

    print(
        "QUESTION:",
        q[3]
    )

    print("====================================\n")


    # ======================================================
    # CALCULATE REMAINING TIME
    # ======================================================

    remaining = (
        exam_duration
        -
        (
            time.time()
            -
            session.get(
                "start_time",
                time.time()
            )
        )
    )


    # ======================================================
    # DISPLAY QUESTION
    # ======================================================

    return render_template(
        "question.html",

        q=q,

        remaining=max(
            0,
            int(remaining)
        ),

        total_questions=len(
            questions
        ),

        current_question=(
            current_index + 1
        )
    )
    


# ==========================================================
# CHECK ANSWER
#
# V2 QUESTION BANK FOR JAMB / WAEC
# POST-UTME REMAINS ON CURRENT SYSTEM
# ==========================================================

@app.route("/check_answer", methods=["POST"])
def check_answer():

    # ======================================================
    # GET SELECTED ANSWER
    # ======================================================

    selected = request.form.get(
        "answer",
        ""
    ).strip().upper()

    question_id = session.get(
        "current_question_id"
    )

    if not question_id:
        return redirect("/result")


    # ======================================================
    # GET EXAM INFORMATION
    # ======================================================

    exam_type = session.get(
        "exam_type"
    )

    subject = session.get(
        "subject"
    )

    username = session.get(
        "user"
    )

    table_name = session.get(
        "question_source",
        "questions"
    )


    # ======================================================
    # CONNECT DATABASE
    # ======================================================

    conn = sqlite3.connect(
        "database.db"
    )

    cursor = conn.cursor()


    # ======================================================
    # ONLY ALLOW KNOWN QUESTION TABLES
    # ======================================================

    if table_name not in (
        "questions_v2",
        "post_utme_questions"
    ):

        conn.close()

        return "Invalid question source."


    # ======================================================
    # GET CURRENT QUESTION
    # ======================================================

    cursor.execute(
        f"""
        SELECT *
        FROM {table_name}
        WHERE id=?
        """,
        (
            question_id,
        )
    )

    current_q = cursor.fetchone()


    if not current_q:

        conn.close()

        return redirect("/result")


    # ======================================================
    # QUESTIONS_V2
    #
    # IMPORTANT:
    #
    # 7  = question_text
    # 8  = option_a
    # 9  = option_b
    # 10 = option_c
    # 11 = option_d
    # 12 = correct_answer
    # 13 = explanation
    # ======================================================

    if table_name == "questions_v2":

        question_text = current_q[7]

        correct_answer = str(
            current_q[12]
        ).strip().upper()

        explanation = (
            current_q[13]
            or
            "No explanation is available for this question yet."
        )


    # ======================================================
    # POST-UTME
    #
    # KEEP ITS EXISTING COLUMN STRUCTURE
    # ======================================================

    else:

        question_text = current_q[3]

        correct_answer = str(
            current_q[8]
        ).strip().upper()

        explanation = (
            current_q[9]
            or
            "No explanation is available for this question yet."
        )


    # ======================================================
    # DEBUG
    # ======================================================

    print("\n")
    print("================================")
    print("CHECK ANSWER")
    print("QUESTION ID:", question_id)
    print("QUESTION:", question_text)
    print("SUBJECT:", subject)
    print("SELECTED:", selected)
    print("CORRECT ANSWER:", correct_answer)
    print("SOURCE:", table_name)
    print("EXPLANATION:", explanation)
    print("================================")
    print("\n")


    # ======================================================
    # DETERMINE RESULT
    # ======================================================

    is_correct = (
        1
        if selected == correct_answer
        else 0
    )


    if is_correct:

        result = "Correct!"

        session["score"] = (
            session.get(
                "score",
                0
            ) + 1
        )

    else:

        result = "Wrong!"


    # ======================================================
    # UPDATE TOTAL ANSWERED
    # ======================================================

    session["total_answered"] = (
        session.get(
            "total_answered",
            0
        ) + 1
    )


    # ======================================================
    # SUBJECT SCORE
    # ======================================================

    subject_scores = session.get(
        "subject_scores",
        {}
    )

    subject_totals = session.get(
        "subject_totals",
        {}
    )


    subject_totals[subject] = (
        subject_totals.get(
            subject,
            0
        ) + 1
    )


    if is_correct:

        subject_scores[subject] = (
            subject_scores.get(
                subject,
                0
            ) + 1
        )


    session["subject_scores"] = (
        subject_scores
    )

    session["subject_totals"] = (
        subject_totals
    )


    # ======================================================
    # SAVE ANSWER FOR REVIEW
    # ======================================================

    cursor.execute(
        """
        INSERT INTO review_answers
        (
            username,
            exam_type,
            subject,
            question_id,
            question_text,
            selected_answer,
            correct_answer,
            explanation,
            is_correct
        )
        VALUES
        (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            username,
            exam_type,
            subject,
            question_id,
            question_text,
            selected,
            correct_answer,
            explanation,
            is_correct
        )
    )


    # ======================================================
    # GET CURRENT QUESTION INDEX
    # ======================================================

    current_q_index = session.get(
        "q_index",
        0
    )


    # ======================================================
    # UPDATE EXAM PROGRESS
    # ======================================================

    if exam_type in (
        "JAMB",
        "WAEC"
    ):

        cursor.execute(
            """
            UPDATE exam_progress

            SET
                q_index=?,
                score=?,
                total_answered=?,
                updated_at=CURRENT_TIMESTAMP

            WHERE username=?
            AND exam_type=?
            AND subject=?
            AND status='IN_PROGRESS'
            """,
            (
                current_q_index,

                session.get(
                    "score",
                    0
                ),

                session.get(
                    "total_answered",
                    0
                ),

                username,
                exam_type,
                subject
            )
        )


    # ======================================================
    # SAVE DATABASE CHANGES
    # ======================================================

    conn.commit()

    conn.close()


    # ======================================================
    # DISPLAY ANSWER PAGE
    # ======================================================

    return render_template(
        "answer.html",

        result=result,

        explanation=explanation,

        score=session.get(
            "score",
            0
        )
    )
    
    
    
# ==========================================================
# NEXT QUESTION
#
# V2 QUESTION BANK FOR JAMB / WAEC
# POST-UTME REMAINS ON CURRENT SYSTEM
# ==========================================================

@app.route("/next_question")
def next_question():

    # ======================================================
    # CHECK LOGIN
    # ======================================================

    if "user" not in session:
        return redirect("/login")

    # ======================================================
    # GET CURRENT EXAM INFORMATION
    # ======================================================

    exam_type = session.get("exam_type")
    current_subject = session.get("subject")
    username = session.get("user")

    current_index = session.get(
        "q_index",
        0
    )

    # ======================================================
    # DEBUG
    # ======================================================

    print("\n====================================")
    print("NEXT QUESTION")
    print("EXAM TYPE:", exam_type)
    print("CURRENT SUBJECT:", current_subject)
    print("Q_INDEX BEFORE:", current_index)
    print("SUBJECT INDEX:", session.get("subject_index"))
    print("====================================\n")

    # ======================================================
    # GET CURRENT SUBJECT QUESTIONS
    # ======================================================

    if exam_type == "POST-UTME":

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT *
            FROM post_utme_questions
            WHERE university_name=?
            AND subject=?
            ORDER BY id
            """,
            (
                session["university"],
                current_subject
            )
        )

        questions = cursor.fetchall()

        conn.close()

    else:

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT q.id

            FROM questions_v2 q

            INNER JOIN exam_types e
                ON q.exam_type_id = e.id

            INNER JOIN subjects s
                ON q.subject_id = s.id

            WHERE
                e.exam_name=?
                AND s.subject_name=?
                AND s.exam_type_id=e.id
                AND e.status='Active'
                AND s.status='Active'
                AND q.status='Active'

            ORDER BY q.id
            """,
            (
                exam_type,
                current_subject
            )
        )

        questions = cursor.fetchall()

        conn.close()

    # ======================================================
    # CHECK QUESTIONS
    # ======================================================

    if not questions:

        print("\n====================================")
        print("NO QUESTIONS AVAILABLE")
        print("EXAM:", exam_type)
        print("SUBJECT:", current_subject)
        print("====================================\n")

        return "No questions available."

    total_questions = len(questions)

    # ======================================================
    # MOVE TO NEXT QUESTION
    # ======================================================

    next_index = current_index + 1

    # ======================================================
    # CURRENT SUBJECT STILL HAS QUESTIONS
    # ======================================================

    if next_index < total_questions:

        session["q_index"] = next_index

        session.modified = True

        # --------------------------------------------------
        # SAVE PROGRESS
        # --------------------------------------------------

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE exam_progress
            SET
                q_index=?,
                score=?,
                total_answered=?,
                updated_at=CURRENT_TIMESTAMP
            WHERE username=?
            AND exam_type=?
            AND subject=?
            AND status='IN_PROGRESS'
            """,
            (
                next_index,
                session.get("score", 0),
                session.get("total_answered", 0),
                username,
                exam_type,
                current_subject
            )
        )

        conn.commit()
        conn.close()

        # --------------------------------------------------
        # DEBUG
        # --------------------------------------------------

        print("\n====================================")
        print("NEXT QUESTION")
        print("SUBJECT:", current_subject)
        print("NEW Q_INDEX:", next_index)
        print("QUESTION NUMBER:", next_index + 1)
        print("TOTAL QUESTIONS:", total_questions)
        print("====================================\n")

        return redirect(
            f"/question/{exam_type}/{current_subject}"
        )

    # ======================================================
    # CURRENT SUBJECT IS FINISHED
    # ======================================================

    print("\n====================================")
    print("SUBJECT FINISHED")
    print("EXAM TYPE:", exam_type)
    print("SUBJECT:", current_subject)
    print("FINAL Q_INDEX:", current_index)
    print("TOTAL QUESTIONS:", total_questions)
    print("====================================\n")

    # ======================================================
    # WAEC
    #
    # Keep existing behavior for now.
    # ======================================================

    if exam_type == "WAEC":

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE exam_progress
            SET
                q_index=?,
                score=?,
                total_answered=?,
                status='COMPLETED',
                updated_at=CURRENT_TIMESTAMP
            WHERE username=?
            AND exam_type='WAEC'
            AND subject=?
            AND status='IN_PROGRESS'
            """,
            (
                total_questions,
                session.get("score", 0),
                session.get("total_answered", 0),
                username,
                current_subject
            )
        )

        conn.commit()
        conn.close()

        return redirect("/result")

    # ======================================================
    # JAMB
    # ======================================================

    if exam_type == "JAMB":

        subjects = session.get(
            "subjects",
            []
        )

        current_subject_index = session.get(
            "subject_index",
            0
        )

        # --------------------------------------------------
        # MARK CURRENT SUBJECT AS COMPLETED
        # --------------------------------------------------

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE exam_progress
            SET
                q_index=?,
                score=?,
                total_answered=?,
                current_subject_index=?,
                status='COMPLETED',
                updated_at=CURRENT_TIMESTAMP
            WHERE username=?
            AND exam_type='JAMB'
            AND subject=?
            AND status='IN_PROGRESS'
            """,
            (
                total_questions,
                session.get("score", 0),
                session.get("total_answered", 0),
                current_subject_index,
                username,
                current_subject
            )
        )

        conn.commit()
        conn.close()

        # --------------------------------------------------
        # CALCULATE NEXT SUBJECT
        # --------------------------------------------------

        next_subject_index = (
            current_subject_index + 1
        )

        # --------------------------------------------------
        # ALL JAMB SUBJECTS FINISHED
        # --------------------------------------------------

        if next_subject_index >= len(subjects):

            print("\n====================================")
            print("ALL JAMB SUBJECTS FINISHED")
            print("====================================\n")

            session["subject_index"] = (
                next_subject_index
            )

            session.modified = True

            return redirect("/result")

        # --------------------------------------------------
        # GET NEXT SUBJECT
        # --------------------------------------------------

        next_subject = subjects[
            next_subject_index
        ]

        # --------------------------------------------------
        # UPDATE SESSION
        # --------------------------------------------------

        session["subject_index"] = (
            next_subject_index
        )

        session["subject"] = (
            next_subject
        )

        session["q_index"] = 0

        session.modified = True

        # ==================================================
        # CREATE NEXT SUBJECT PROGRESS
        # ==================================================

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT id
            FROM exam_progress
            WHERE username=?
            AND exam_type='JAMB'
            AND subject=?
            AND status='IN_PROGRESS'
            LIMIT 1
            """,
            (
                username,
                next_subject
            )
        )

        existing = cursor.fetchone()

        # --------------------------------------------------
        # EXISTING IN-PROGRESS RECORD
        # --------------------------------------------------

        if existing:

            cursor.execute(
                """
                UPDATE exam_progress
                SET
                    q_index=0,
                    score=?,
                    total_answered=?,
                    start_time=?,
                    exam_duration=?,
                    current_subject_index=?,
                    updated_at=CURRENT_TIMESTAMP
                WHERE id=?
                """,
                (
                    session.get("score", 0),
                    session.get("total_answered", 0),
                    session.get("start_time"),
                    session.get("exam_duration"),
                    next_subject_index,
                    existing[0]
                )
            )

        # --------------------------------------------------
        # CREATE NEW PROGRESS RECORD
        # --------------------------------------------------

        else:

            cursor.execute(
                """
                INSERT INTO exam_progress
                (
                    username,
                    exam_type,
                    subject,
                    q_index,
                    score,
                    total_answered,
                    start_time,
                    exam_duration,
                    status,
                    current_subject_index
                )
                VALUES
                (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    username,
                    "JAMB",
                    next_subject,
                    0,
                    session.get("score", 0),
                    session.get("total_answered", 0),
                    session.get("start_time"),
                    session.get("exam_duration"),
                    "IN_PROGRESS",
                    next_subject_index
                )
            )

        conn.commit()
        conn.close()

        # --------------------------------------------------
        # DEBUG
        # --------------------------------------------------

        print("\n====================================")
        print("MOVING TO NEXT JAMB SUBJECT")
        print("CURRENT SUBJECT:", current_subject)
        print("NEXT SUBJECT:", next_subject)
        print("NEXT SUBJECT INDEX:", next_subject_index)
        print("Q_INDEX RESET TO:", 0)
        print("SCORE:", session.get("score", 0))
        print("TOTAL ANSWERED:", session.get("total_answered", 0))
        print("====================================\n")

        # --------------------------------------------------
        # GO TO NEXT SUBJECT
        # --------------------------------------------------

        return redirect(
            f"/question/JAMB/{next_subject}"
        )

    # ======================================================
    # SAFETY FALLBACK
    # ======================================================

    return redirect("/result")
    
    
    
# ✅ Manage Questions with Search + Filters + Pagination
@app.route("/manage_questions")
def manage_questions():

    if "admin" not in session:
        return redirect("/admin_login")

    search = request.args.get("search", "").strip()

    exam_filter = request.args.get("exam_type", "").strip()

    subject_filter = request.args.get("subject", "").strip()

    page = request.args.get("page", 1, type=int)

    per_page = 20

    offset = (page - 1) * per_page

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    conditions = []
    params = []

    if search:
        conditions.append(
            "(exam_type LIKE ? OR subject LIKE ? OR question_text LIKE ?)"
        )

        params.extend([
            f"%{search}%",
            f"%{search}%",
            f"%{search}%"
        ])

    if exam_filter:
        conditions.append("exam_type = ?")
        params.append(exam_filter)

    if subject_filter:
        conditions.append("subject = ?")
        params.append(subject_filter)

    where_clause = ""

    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    # Count total records
    count_query = f"""
        SELECT COUNT(*)
        FROM questions
        {where_clause}
    """

    cursor.execute(count_query, params)

    total_questions = cursor.fetchone()[0]

    # Get paginated records
    query = f"""
        SELECT *
        FROM questions
        {where_clause}
        ORDER BY id DESC
        LIMIT ?
        OFFSET ?
    """

    cursor.execute(
        query,
        params + [per_page, offset]
    )

    questions = cursor.fetchall()

    conn.close()

    total_pages = (
        total_questions + per_page - 1
    ) // per_page

    return render_template(
        "manage_questions.html",
        questions=questions,
        search=search,
        exam_filter=exam_filter,
        subject_filter=subject_filter,
        page=page,
        total_pages=total_pages
    )
    
# ✅ Edit Question
@app.route("/edit_question/<int:question_id>", methods=["GET", "POST"])
def edit_question(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Save updates
    if request.method == "POST":

        exam_type = request.form["exam_type"]
        subject = request.form["subject"]
        question_text = request.form["question"]
        option_a = request.form["option_a"]
        option_b = request.form["option_b"]
        option_c = request.form["option_c"]
        option_d = request.form["option_d"]
        correct_answer = request.form["correct_answer"]
        explanation = request.form["explanation"]

        cursor.execute(
            """
            UPDATE questions
            SET
                exam_type=?,
                subject=?,
                question_text=?,
                option_a=?,
                option_b=?,
                option_c=?,
                option_d=?,
                correct_answer=?,
                explanation=?
            WHERE id=?
            """,
            (
                exam_type,
                subject,
                question_text,
                option_a,
                option_b,
                option_c,
                option_d,
                correct_answer,
                explanation,
                question_id
            )
        )

        conn.commit()
        conn.close()

        return redirect("/manage_questions")

    # Load question details
    cursor.execute(
        """
        SELECT *
        FROM questions
        WHERE id=?
        """,
        (question_id,)
    )

    question = cursor.fetchone()

    conn.close()

    if not question:
        return "Question not found."

    return render_template(
        "edit_question.html",
        question=question
    )
    
# ✅ Delete Question
@app.route("/delete_question/<int:question_id>", methods=["GET", "POST"])
def delete_question(question_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get question details
    cursor.execute(
        """
        SELECT *
        FROM questions
        WHERE id=?
        """,
        (question_id,)
    )

    question = cursor.fetchone()

    if not question:
        conn.close()
        return "Question not found."

    # Delete after confirmation
    if request.method == "POST":

        cursor.execute(
            """
            DELETE FROM questions
            WHERE id=?
            """,
            (question_id,)
        )

        conn.commit()
        conn.close()

        return redirect("/manage_questions")

    conn.close()

    return render_template(
        "delete_question.html",
        question=question
    )
    
# ✅ Export Questions to Excel
@app.route("/export_questions")
def export_questions():

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")

    query = """
        SELECT
            exam_type,
            subject,
            question_text,
            option_a,
            option_b,
            option_c,
            option_d,
            correct_answer,
            explanation
        FROM questions
        ORDER BY id DESC
    """

    df = pd.read_sql_query(query, conn)

    conn.close()

    filename = "cbt_questions_export.xlsx"

    df.to_excel(
        filename,
        index=False
    )

    return send_file(
        filename,
        as_attachment=True
    )    


# ✅ Result page
@app.route("/result")
def result():

    score = session.get("score", 0)

    total_questions = session.get("total_answered", 0)

    wrong_answers = total_questions - score

    # prevent division by zero
    if total_questions == 0:
        percentage = 0
    else:
        percentage = int((score / total_questions) * 100)

    # performance message
    if percentage >= 70:
        performance = "Excellent Performance 🎉"
        status = "PASS"

    elif percentage >= 50:
        performance = "Good Job 👍"
        status = "AVERAGE"

    else:
        performance = "Needs Improvement 📚"
        status = "FAIL"

    # Save result to database
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    exam_type = session.get("exam_type")

    if exam_type == "JAMB":
        exam_name = session.get("course")
    else:
        exam_name = session.get("subject")


    duration = int(
        time.time() - session.get("start_time", time.time())
    )


    # Generate verification code
    import datetime

    verification_code = (
        f"CBT-{datetime.datetime.now().strftime('%Y%m%d')}-"
        f"{int(time.time())}"
    )

    cursor.execute(
        """
        INSERT INTO results
        (
            username,
            exam_type,
            score,
            total,
            percentage,
            date_taken,
            exam_name,
            duration,
            status,
            verification_code
        )
        VALUES (?, ?, ?, ?, ?, datetime('now'), ?, ?, ?, ?)
        """,
        (
            session.get("user"),
            exam_type,
            score,
            total_questions,
            percentage,
            exam_name,
            duration,
            status,
            verification_code
        )
    )

    conn.commit()
    
    # ✅ Mark unfinished exam as completed
    cursor.execute(
        """
        UPDATE exam_progress
        SET
            status='COMPLETED',
            updated_at=CURRENT_TIMESTAMP
        WHERE username=?
        AND status='IN_PROGRESS'
        """,
        (
            session.get("user"),
        )
    )

    conn.commit()
    
    # Get the ID of the newly saved result
    result_id = cursor.lastrowid
    
    subject_scores = session.get("subject_scores", {})
    subject_totals = session.get("subject_totals", {})
    
    
    for subject, score in subject_scores.items():

        total = subject_totals.get(subject, 0)

        if total > 0:
            percentage = round((score / total) * 100, 2)
        else:
            percentage = 0

        cursor.execute(
            """
            INSERT INTO result_subjects
            (
                result_id,
                username,
                exam_type,
                exam_name,
                subject,
                score,
                total_questions,
                percentage,
                time_spent
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                result_id,
                session.get("user"),
                exam_type,
                exam_name,
                subject,
                score,
                total,
                percentage,
                duration
            )
        )

    conn.commit()

        
    # Save review answers permanently
    review_answers = session.get("review_answers", [])

    for answer in review_answers:

        cursor.execute(
            """
            INSERT INTO review_answers
            (
                result_id,
                username,
                subject,
                question_text,
                selected_answer,
                correct_answer,
                explanation,
                is_correct
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                result_id,
                session.get("user"),
                answer["subject"],
                answer["question_text"],
                answer["selected_answer"],
                answer["correct_answer"],
                answer["explanation"],
                answer["is_correct"]
            )
        )
        

    conn.commit()
    conn.close()

    # reset exam session
    session.pop("q_index", None)
    session.pop("score", None)
    session.pop("subject_index", None)
    session.pop("subjects", None)
    session.pop("start_time", None)
    session.pop("exam_duration", None)
    session.pop("total_answered", None)
    session.pop("current_question_id", None)
    session.pop("review_answers", None)

    for key in list(session.keys()):
        if key.startswith("shuffled_"):
            session.pop(key, None)

    # add these
    session.pop("course", None)
    session.pop("exam_type", None)
    session.pop("subject", None)
    
    subject_scores = session.get("subject_scores", {})
    subject_totals = session.get("subject_totals", {})

    return render_template(
        "result.html",
        score=score,
        total=total_questions,
        wrong_answers=wrong_answers,
        percentage=percentage,
        performance=performance,
        status=status,
        subject_scores=subject_scores,
        subject_totals=subject_totals
    )

# ✅ Professional My Results Page
@app.route("/my_results")
def my_results():

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM results
        WHERE username=?
        ORDER BY id DESC
        """,
        (session["user"],)
    )

    rows = cursor.fetchall()

    conn.close()

    formatted_results = []

    for r in rows:

        duration = r[8]

        # Format duration nicely
        if duration is None:
            duration_text = "-"

        elif duration < 60:
            duration_text = f"{duration} sec"

        elif duration < 3600:
            minutes = duration // 60
            seconds = duration % 60
            duration_text = f"{minutes} min {seconds} sec"

        else:
            hours = duration // 3600
            minutes = (duration % 3600) // 60
            duration_text = f"{hours} hr {minutes} min"

        formatted_results.append({
            
             "id": r[0],

            "date_taken": r[6],

            "exam_type": r[2],

            # Handle old records gracefully
            "exam_name": r[7] if r[7] else "-",

            "score": r[3],

            "total": r[4],

            "percentage": r[5],

            "duration": duration_text,

            "status": r[9] if r[9] else "-"

        })

    return render_template(
        "my_results.html",
        results=formatted_results
    )

# ✅ Verify Result
@app.route("/verify_result/<code>")
def verify_result(code):

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            results.*,
            users.name,
            users.phone
        FROM results
        LEFT JOIN users
        ON results.username = users.email
        WHERE results.verification_code=?
        """,
        (code,)
    )

    result = cursor.fetchone()

    conn.close()

    if not result:
        return render_template(
            "verify_result.html",
            verified=False
        )

    return render_template(
        "verify_result.html",
        verified=True,
        result=result
    )

# ✅ Download Professional PDF Result Slip
@app.route("/download_result/<int:result_id>")
def download_result(result_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    # Get result details
    cursor.execute(
        """
        SELECT
            r.exam_type,
            r.exam_name,
            r.score,
            r.total,
            r.percentage,
            r.date_taken,
            r.duration,
            r.status,
            r.verification_code,
            u.name,
            u.email,
            u.phone
        FROM results r
        JOIN users u
            ON r.username = u.email
        WHERE r.id = ?
        """,
        (
            (result_id,)
        )
    )

    result = cursor.fetchone()

    conn.close()

    if not result:
        return "Result not found."

    (
        exam_type,
        exam_name,
        score,
        total,
        percentage,
        date_taken,
        duration,
        status,
        verification_code,
        student_name,
        student_email,
        student_phone
    ) = result
    
    # Generate QR only if verification code exists

    qr_buffer = None

    if verification_code:

        verification_url = (
            request.host_url.rstrip("/")
            + "/verify_result/"
            + str(verification_code)
        )

        qr = qrcode.make(verification_url)

        qr_buffer = BytesIO()

        qr.save(qr_buffer, format="PNG")

        qr_buffer.seek(0)

    # Format duration
    if duration is None:
        duration_text = "-"

    elif duration < 60:
        duration_text = f"{duration} sec"

    elif duration < 3600:
        minutes = duration // 60
        seconds = duration % 60
        duration_text = f"{minutes} min {seconds} sec"

    else:
        hours = duration // 3600
        minutes = (duration % 3600) // 60
        duration_text = f"{hours} hr {minutes} min"

    # Create PDF in memory
    buffer = BytesIO()

    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()

    elements = []

    # System Name
    elements.append(
        Paragraph(
            "<font color='blue'><b>NIGERIA CBT EXAMINATION SYSTEM</b></font>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Official Student Result Slip",
            styles["Heading2"]
        )
    )

    elements.append(Spacer(1, 20))

    student_data = [

    ["Student Name", student_name],
    ["Email", student_email],
    ["Phone", student_phone or "-"]

    ]

    student_table = Table(
        student_data,
        colWidths=[2.2*inch, 4*inch]
    )

    student_table.setStyle(TableStyle([

        ('BACKGROUND', (0,0), (0,-1), colors.lightblue),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('PADDING', (0,0), (-1,-1), 8)

    ]))

    elements.append(student_table)

    elements.append(Spacer(1, 20))

    elements.append(Spacer(1, 15))

    exam_data = [

    ["Exam Type", exam_type],
    ["Subject/Course", exam_name],
    ["Score", f"{score}/{total}"],
    ["Percentage", f"{percentage}%"],
    ["Status", status],
    ["Time Spent", duration_text],
    ["Date Taken", date_taken],
    ["Verification Code", verification_code]

    ]

    exam_table = Table(
        exam_data,
        colWidths=[2.2*inch, 4*inch]
    )

    exam_table.setStyle(TableStyle([

        ('BACKGROUND', (0,0), (0,-1), colors.beige),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('PADDING', (0,0), (-1,-1), 8)

    ]))

    elements.append(exam_table)
    
    elements.append(Spacer(1, 20))

    if qr_buffer:

        elements.append(
            Paragraph(
                "<b>Scan QR Code to Verify Result</b>",
                styles["Heading3"]
            )
        )

        qr_image = Image(qr_buffer)

        qr_image.drawWidth = 1.5 * inch
        qr_image.drawHeight = 1.5 * inch

        elements.append(qr_image)

    elements.append(Spacer(1, 25))

    elements.append(Spacer(1, 30))

    elements.append(
        Paragraph(
            "<i>This is an electronically generated result slip.</i>",
            styles["Italic"]
        )
    )

    elements.append(
        Paragraph(
            "<i>No signature is required.</i>",
            styles["Italic"]
        )
    )

    doc.build(elements)

    buffer.seek(0)

    filename = (
        f"{exam_type}_{exam_name}_Result.pdf"
    )

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf"
    )
    
# ✅ Result Details Page
@app.route("/result_details/<int:result_id>")
def result_details(result_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM results
        WHERE id=? AND username=?
        """,
        (
            result_id,
            session["user"]
        )
    )

    result = cursor.fetchone()
    
    cursor.execute(
        """
        SELECT
            subject,
            score,
            total_questions,
            percentage,
            time_spent
        FROM result_subjects
        WHERE result_id=?
        ORDER BY id
        """,
        (result_id,)
    )

    subject_results = cursor.fetchall()

    conn.close()

    if not result:
        return "Result not found."

    duration = result[8]

    if duration is None:
        duration_text = "-"

    elif duration < 60:
        duration_text = f"{duration} sec"

    elif duration < 3600:
        minutes = duration // 60
        seconds = duration % 60
        duration_text = f"{minutes} min {seconds} sec"

    else:
        hours = duration // 3600
        minutes = (duration % 3600) // 60
        duration_text = f"{hours} hr {minutes} min"

    return render_template(
        "result_details.html",

        result={
            "id": result_id,
            "date_taken": result[6],
            "exam_type": result[2],
            "exam_name": result[7] if result[7] else "-",
            "score": result[3],
            "total": result[4],
            "percentage": result[5],
            "duration": duration_text,
            "status": result[9] if result[9] else "-",
            "verification_code": result[10] if result[10] else "-"
        },

        subject_results=subject_results,
        student_name=session.get("name", session.get("user")),
        student_email=session.get("user")
    )


# ✅ Review Answers
@app.route("/review_answers/<int:result_id>")
def review_answers(result_id):

    if "user" not in session:
        return redirect("/login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            subject,
            question_text,
            selected_answer,
            correct_answer,
            explanation,
            is_correct
        FROM review_answers
        WHERE result_id=? AND username=?
        """,
        (
            result_id,
            session["user"]
        )
    )

    answers = cursor.fetchall()

    conn.close()

    return render_template(
        "review_answers.html",
        answers=answers
    )
    
    
# ✅ Admin Review Answers
@app.route("/admin/review_answers/<int:result_id>")
def admin_review_answers(result_id):

    if "admin" not in session:
        return redirect("/admin_login")

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            subject,
            question_text,
            selected_answer,
            correct_answer,
            explanation,
            is_correct
        FROM review_answers
        WHERE result_id=?
    """, (result_id,))

    answers = cursor.fetchall()

    conn.close()

    return render_template(
        "review_answers.html",
        answers=answers
    )
    
    
    
# ✅ Bulk Upload Questions
@app.route("/upload_questions", methods=["GET", "POST"])
def upload_questions():

    if "admin" not in session:
        return redirect("/admin_login")

    if request.method == "POST":

        file = request.files.get("file")

        if not file:
            return "No file selected."

        try:

            # Read file
            if file.filename.endswith(".csv"):
                df = pd.read_csv(file)

            elif file.filename.endswith(".xlsx"):
                df = pd.read_excel(file)

            else:
                return "Only CSV and XLSX files are allowed."

            conn = sqlite3.connect("database.db")
            cursor = conn.cursor()

            imported = 0
            duplicates = 0
            errors = []

            for index, row in df.iterrows():

                try:

                    exam_type = str(row["exam_type"]).strip()
                    subject = str(row["subject"]).strip()
                    question_text = str(row["question_text"]).strip()

                    # Validate required fields
                    if not exam_type:
                        raise ValueError("Exam type is missing")

                    if not subject:
                        raise ValueError("Subject is missing")

                    if not question_text:
                        raise ValueError("Question text is missing")

                    correct_answer = str(
                        row["correct_answer"]
                    ).strip().upper()

                    if correct_answer not in ["A", "B", "C", "D"]:
                        raise ValueError(
                            "Correct answer must be A, B, C or D"
                        )

                    # Check if question already exists
                    cursor.execute(
                        """
                        SELECT id
                        FROM questions
                        WHERE exam_type=?
                        AND subject=?
                        AND question_text=?
                        """,
                        (
                            exam_type,
                            subject,
                            question_text
                        )
                    )

                    existing = cursor.fetchone()

                    if existing:
                        duplicates += 1
                        continue

                    # Insert question
                    cursor.execute(
                        """
                        INSERT INTO questions
                        (
                            exam_type,
                            subject,
                            question_text,
                            option_a,
                            option_b,
                            option_c,
                            option_d,
                            correct_answer,
                            explanation
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            exam_type,
                            subject,
                            question_text,
                            str(row["option_a"]).strip(),
                            str(row["option_b"]).strip(),
                            str(row["option_c"]).strip(),
                            str(row["option_d"]).strip(),
                            correct_answer,
                            str(row["explanation"]).strip()
                        )
                    )

                    imported += 1

                except Exception as e:

                    errors.append(
                        f"Row {index + 2}: {str(e)}"
                    )

            conn.commit()
            conn.close()

            error_text = "<br>".join(errors)

            return f"""
            ✅ {imported} questions imported successfully.<br>
            ⚠️ {duplicates} duplicate questions skipped.<br>
            ❌ {len(errors)} rows had errors.<br><br>
            {error_text}
            """

        except Exception as e:

            return f"Error: {e}"

    return render_template("upload_questions.html")


# ==========================================================
# RESUME UNFINISHED EXAM
# ==========================================================

@app.route("/resume_exam")
def resume_exam():

    # ======================================================
    # CHECK LOGIN
    # ======================================================

    if "user" not in session:
        return redirect("/login")

    username = session["user"]

    # ======================================================
    # GET UNFINISHED EXAM FROM DATABASE
    # ======================================================

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT
            id,
            exam_type,
            subject,
            q_index,
            score,
            total_answered,
            start_time,
            exam_duration,
            current_subject_index
        FROM exam_progress
        WHERE username=?
        AND status='IN_PROGRESS'
        ORDER BY id DESC
        LIMIT 1
        """,
        (username,)
    )

    progress = cursor.fetchone()

    conn.close()

    # ======================================================
    # NO UNFINISHED EXAM
    # ======================================================

    if not progress:

        return redirect("/dashboard")

    # ======================================================
    # UNPACK DATABASE PROGRESS
    # ======================================================

    (
        progress_id,
        exam_type,
        saved_subject,
        saved_q_index,
        saved_score,
        saved_total_answered,
        saved_start_time,
        saved_exam_duration,
        saved_subject_index
    ) = progress

    # ======================================================
    # JAMB
    # ======================================================

    if exam_type == "JAMB":

        # --------------------------------------------------
        # FIND THE COURSE THAT CONTAINS THE SAVED SUBJECT
        # --------------------------------------------------

        saved_course = None

        for course_name, course_subjects in JAMB_COURSES.items():

            if saved_subject in course_subjects:

                saved_course = course_name

                break

        # --------------------------------------------------
        # SAFETY CHECK
        # --------------------------------------------------

        if not saved_course:

            return (
                "Unable to determine the JAMB course "
                "for the unfinished examination."
            )

        subjects = JAMB_COURSES.get(
            saved_course
        )

        # --------------------------------------------------
        # RESTORE COMPLETE JAMB SESSION
        # --------------------------------------------------

        session["exam_type"] = "JAMB"

        session["course"] = saved_course

        session["subjects"] = subjects

        session["subject_index"] = (
            saved_subject_index or 0
        )

        session["subject"] = saved_subject

        session["q_index"] = (
            saved_q_index or 0
        )

        session["score"] = (
            saved_score or 0
        )

        session["total_answered"] = (
            saved_total_answered or 0
        )

        session["start_time"] = (
            saved_start_time
        )

        session["exam_duration"] = (
            saved_exam_duration or (120 * 60)
        )

        # --------------------------------------------------
        # Restore subject score dictionaries
        # --------------------------------------------------

        session.setdefault(
            "subject_scores",
            {}
        )

        session.setdefault(
            "subject_totals",
            {}
        )

        # --------------------------------------------------
        # Clear temporary question information
        # --------------------------------------------------

        session.pop(
            "current_question_id",
            None
        )

        session.pop(
            "question_source",
            None
        )

        # --------------------------------------------------
        # Clear temporary shuffled data
        # --------------------------------------------------

        for key in list(session.keys()):

            if key.startswith("shuffled_"):

                session.pop(
                    key,
                    None
                )

        # --------------------------------------------------
        # DEBUG
        # --------------------------------------------------

        print("\n====================================")
        print("RESUMING JAMB EXAM")
        print("PROGRESS ID:", progress_id)
        print("COURSE:", saved_course)
        print("SUBJECT:", saved_subject)
        print("SUBJECT INDEX:", saved_subject_index)
        print("QUESTION INDEX:", saved_q_index)
        print("SCORE:", saved_score)
        print("TOTAL ANSWERED:", saved_total_answered)
        print("====================================\n")

    # ======================================================
    # WAEC / OTHER EXAM TYPES
    # ======================================================

    else:

        session["exam_type"] = exam_type

        session["subject"] = saved_subject

        session["q_index"] = (
            saved_q_index or 0
        )

        session["score"] = (
            saved_score or 0
        )

        session["total_answered"] = (
            saved_total_answered or 0
        )

        session["start_time"] = (
            saved_start_time
        )

        session["exam_duration"] = (
            saved_exam_duration or 120 * 60
        )

    # ======================================================
    # RESUME EXACT QUESTION
    # ======================================================

    return redirect(
        f"/question/{exam_type}/{saved_subject}"
    )


# ==========================================================
# DISCARD UNFINISHED EXAM AND START A NEW ONE
# ==========================================================

@app.route("/start_new_exam")
def start_new_exam():

    # ======================================================
    # CHECK LOGIN
    # ======================================================

    if "user" not in session:
        return redirect("/login")

    username = session["user"]

    # ======================================================
    # DELETE ALL UNFINISHED EXAM PROGRESS
    # ======================================================

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute(
        """
        DELETE FROM exam_progress
        WHERE username=?
        AND status='IN_PROGRESS'
        """,
        (
            username,
        )
    )

    conn.commit()
    conn.close()

    # ======================================================
    # CLEAR EXAM SESSION DATA
    # ======================================================

    session.pop("exam_type", None)
    session.pop("course", None)
    session.pop("subjects", None)
    session.pop("subject_index", None)
    session.pop("subject", None)
    session.pop("q_index", None)
    session.pop("score", None)
    session.pop("total_answered", None)
    session.pop("subject_scores", None)
    session.pop("subject_totals", None)
    session.pop("start_time", None)
    session.pop("exam_duration", None)
    session.pop("current_question_id", None)
    session.pop("question_source", None)
    session.pop("review_answers", None)

    # ======================================================
    # CLEAR SHUFFLED QUESTION DATA
    # ======================================================

    for key in list(session.keys()):

        if key.startswith("shuffled_"):

            session.pop(
                key,
                None
            )

    # ======================================================
    # GO BACK TO EXAM TYPES
    # ======================================================

    return redirect("/exam_types")
    

# ✅ Logout
@app.route("/logout")
def logout():

    username = session.get("user")
    session_token = session.get("session_token")

    if username and session_token:

        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()

        cursor.execute(
            """
            UPDATE user_sessions
            SET is_active = 0
            WHERE username = ?
            AND session_token = ?
            """,
            (
                username,
                session_token
            )
        )

        conn.commit()
        conn.close()

    # Completely clear the Flask session
    session.clear()

    return redirect("/?logged_out=1")



# ==========================================
# CREATE QUESTIONS_V2 TABLE
# ==========================================

def create_questions_v2_table():

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS questions_v2 (

            id INTEGER PRIMARY KEY AUTOINCREMENT,

            exam_type_id INTEGER NOT NULL,

            subject_id INTEGER NOT NULL,

            topic_id INTEGER NOT NULL,

            passage_id INTEGER,

            question_type TEXT DEFAULT 'Objective',

            difficulty TEXT DEFAULT 'Medium',

            question_text TEXT NOT NULL,

            option_a TEXT,

            option_b TEXT,

            option_c TEXT,

            option_d TEXT,

            correct_answer TEXT NOT NULL,

            explanation TEXT,

            status TEXT DEFAULT 'Active',

            created_at TEXT DEFAULT CURRENT_TIMESTAMP,

            FOREIGN KEY (exam_type_id)
                REFERENCES exam_types(id),

            FOREIGN KEY (subject_id)
                REFERENCES subjects(id),

            FOREIGN KEY (topic_id)
                REFERENCES topics(id)

        )
    """)

    conn.commit()
    conn.close()



# ✅ Run app
if __name__ == "__main__":

    init_db()

    create_questions_v2_table()

    app.run(debug=True)