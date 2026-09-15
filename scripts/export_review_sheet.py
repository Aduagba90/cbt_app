"""Export a subject's curated (tier-1) questions to an Excel sheet a teacher can mark.

    python3 scripts/export_review_sheet.py "Use of English" docs/review/english_review.xlsx

One row per question: passage (if any), question, options A-D, our answer, explanation, and an empty
"Verdict" column (OK / Wrong answer / Too easy / Unclear) plus "Comment". The teacher returns the file;
fixes are then written into the next content batch.
"""
import os
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.environ.get("DATABASE_PATH", os.path.join(ROOT, "database.db"))


def export(subject, out_path, exam_type="JAMB"):
    con = sqlite3.connect(DB)
    rows = con.execute(
        """SELECT q.id, t.topic_name, p.title, p.passage_text, q.question_text,
                  q.option_a, q.option_b, q.option_c, q.option_d, q.correct_answer, q.explanation
           FROM questions_v2 q
           JOIN subjects s ON s.id = q.subject_id
           JOIN exam_types e ON e.id = q.exam_type_id
           LEFT JOIN topics t ON t.id = q.topic_id
           LEFT JOIN passages p ON p.id = q.passage_id
           WHERE e.exam_name = ? AND s.subject_name = ? AND q.status = 'Active' AND q.tier = 1
           ORDER BY COALESCE(q.passage_id, 0) DESC, q.passage_id, t.topic_name, q.id""",
        (exam_type, subject),
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    head = ["No", "ID", "Section", "Passage title", "Passage text", "Question", "A", "B", "C", "D",
            "Our answer", "Explanation", "Verdict", "Comment"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    last_passage = None
    for n, r in enumerate(rows, start=1):
        qid, topic, title, ptext, q, a, b, c, d, ans, expl = r
        # Only print the passage text on its first question, so the sheet stays readable.
        show = ptext if (title and title != last_passage) else ""
        last_passage = title
        ws.append([n, qid, topic, title or "", show, q, a, b, c, d, ans, expl or "", "", ""])
    widths = [5, 7, 18, 22, 60, 55, 22, 22, 22, 22, 9, 50, 14, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    dv = DataValidation(type="list", formula1='"OK,Wrong answer,Too easy,Unclear"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{len(rows) + 1}")
    ws.freeze_panes = "F2"
    guide = wb.create_sheet("How to review")
    for line in [
        "Thank you for checking these questions.",
        "",
        "1. Read each question (and its passage, shown on the first question of that passage).",
        "2. In the Verdict column choose one of: OK / Wrong answer / Too easy / Unclear.",
        "3. If it is not OK, write what should change in the Comment column (e.g. the correct letter, or a better wording).",
        "4. Leave the rest as it is and send the file back.",
        "",
        f"Subject: {subject} ({exam_type}). Questions in this file: {len(rows)}.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    con.close()
    return len(rows)


if __name__ == "__main__":
    subj = sys.argv[1] if len(sys.argv) > 1 else "Use of English"
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "docs", "review", subj.lower().replace(" ", "_") + "_review.xlsx")
    n = export(subj, out)
    print(f"wrote {out}: {n} questions")
